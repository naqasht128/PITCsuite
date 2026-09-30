"""
CP136 & CP137 Change Posting Report Engine for PITC Suite
Parses WAPDA RCP-136 & CP-137 change posting reports (.txt and .pdf),
extracts 14-digit Ref Nos, modified fields below or against headings,
handles deduplication (single postings), and exports formatted multi-sheet Excel (.xlsx) files.
"""

import os
import re
import glob
from typing import List, Dict, Any, Tuple, Optional, Callable
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

try:
    import pdfplumber
except ImportError:
    pdfplumber = None


# Ordered list of all standard headings in WAPDA RCP-136 and CP-137 reports
HEADINGS_CORE = [
    'Sub Division', 'Customer ID'
]

HEADINGS_TOP = [
    'SUB BATCH', 'TARIFF', 'APPLICATION NO', 'APP DATE', 'SEC AMOUNT', 'SEC DATE',
    'MAIN SETS', 'SUB SETS', 'S-POINTS'
]

HEADINGS_BLOCK_1 = [
    'CONN DATE', 'SAN LOAD', 'TRANS. CAPCTY', 'RU', 'F/P', 'SEASON CD', 'CONS. CLASS',
    'ED CD', 'GVT DEP', 'STD CLSF', 'IT', 'METERS', 'AC', 'NTN', 'GST'
]

HEADINGS_BLOCK_2 = [
    'METER RENT', 'SRV RENT', 'N .ID NUMBER', 'TRANSFRMR CODE',
    'STREET LIGHT CATEGORY', 'STREET LIGHT LINE LENGTH'
]

HEADINGS_LABELS = [
    'CONTACT #', 'EMAIL ADDRESS', 'CONNECTION CATEGORY CODE',
    'TRANSFORMER REF NO', 'TRANSFORMER ADDRESS',
    'CONSUMER NAME', 'FATHER/HUSBAND NAME',
    'ADDRES LINE -1', 'ADDRES LINE -2',
    'SIM ID', 'SIM NO',
    'NET MTR CD', 'NET MTR CAP',
    'NET MTR CONNECTION DATE', 'NET MTR EXPIRY DATE',
    'LICENSE#'
]

HEADINGS_BLOCK_3 = [
    'SET NO', 'METER TYPE', 'MAIN/S. SET', 'KVARH/KVAH-CD',
    'FEEDER CODE', 'METER NO', 'LOC CODE', 'MFG CODE', 'METER CODE',
    'M/PLY FACTOR', 'MTR RANGE', 'METER READING'
]

HEADINGS_BOTTOM = [
    'NEW-REFERENCE-NUMBER', 'A/CODE', 'BANK NAME', 'AMOUNT', 'START DATE', 'EXPIRY DATE', 'ERRORS CODES'
]

ALL_ORDERED_HEADINGS = (
    HEADINGS_CORE + HEADINGS_TOP + HEADINGS_BLOCK_1 + HEADINGS_BLOCK_2 +
    HEADINGS_LABELS + HEADINGS_BLOCK_3 + HEADINGS_BOTTOM
)

# Fixed-width column spans for text (.txt) files
SPANS_BLOCK_1_CP136: List[Tuple[str, int, int]] = [
    ('CONN DATE', 0, 14),
    ('SAN LOAD', 14, 27),
    ('RU', 27, 34),
    ('F/P', 34, 42),
    ('SEASON CD', 42, 55),
    ('ED CD', 55, 64),
    ('GVT DEP', 64, 73),
    ('STD CLSF', 73, 82),
    ('IT', 82, 87),
    ('METERS', 87, 97),
    ('AC', 97, 103),
    ('NTN', 103, 115),
    ('GST', 115, 135),
]

SPANS_BLOCK_1_CP137: List[Tuple[str, int, int]] = [
    ('CONN DATE', 0, 16),
    ('SAN LOAD', 16, 26),
    ('TRANS. CAPCTY', 26, 38),
    ('RU', 38, 45),
    ('F/P', 45, 53),
    ('SEASON CD', 53, 65),
    ('CONS. CLASS', 65, 76),
    ('ED CD', 76, 86),
    ('GVT DEP', 86, 96),
    ('STD CLSF', 96, 106),
    ('IT', 106, 116),
    ('METERS', 116, 126),
    ('AC', 126, 134),
    ('NTN', 134, 146),
    ('GST', 146, 165),
]

SPANS_BLOCK_1 = SPANS_BLOCK_1_CP136

SPANS_BLOCK_2_CP136: List[Tuple[str, int, int]] = [
    ('METER RENT', 0, 15),
    ('SRV RENT', 15, 28),
    ('N .ID NUMBER', 28, 44),
    ('TRANSFRMR CODE', 44, 60),
]

SPANS_BLOCK_2_CP137: List[Tuple[str, int, int]] = [
    ('METER RENT', 0, 15),
    ('SRV RENT', 15, 28),
    ('N .ID NUMBER', 28, 44),
    ('TRANSFRMR CODE', 44, 60),
    ('STREET LIGHT CATEGORY', 60, 85),
    ('STREET LIGHT LINE LENGTH', 85, 110),
]

SPANS_BLOCK_2 = SPANS_BLOCK_2_CP136

SPANS_BLOCK_3_CP136: List[Tuple[str, int, int]] = [
    ('SET NO', 0, 9),
    ('METER TYPE', 9, 21),
    ('FEEDER CODE', 21, 31),
    ('METER NO', 31, 40),
    ('LOC CODE', 40, 49),
    ('MFG CODE', 49, 58),
    ('METER CODE', 58, 66),
    ('M/PLY FACTOR', 66, 76),
    ('MTR RANGE', 76, 83),
    ('METER READING', 83, 97),
    ('ERRORS CODES', 97, 135),
]

SPANS_BLOCK_3_CP137: List[Tuple[str, int, int]] = [
    ('SET NO', 0, 9),
    ('METER TYPE', 9, 18),
    ('MAIN/S. SET', 18, 28),
    ('KVARH/KVAH-CD', 28, 38),
    ('FEEDER CODE', 38, 48),
    ('METER NO', 48, 58),
    ('LOC CODE', 58, 67),
    ('MFG CODE', 67, 76),
    ('METER CODE', 76, 85),
    ('M/PLY FACTOR', 85, 94),
    ('MTR RANGE', 94, 102),
    ('METER READING', 102, 115),
    ('ERRORS CODES', 115, 140),
]

SPANS_BLOCK_3 = SPANS_BLOCK_3_CP136



def clean_page_banners(raw_text: str) -> List[str]:
    """
    Splits text across form-feed characters and strips standard page header banners.
    """
    pages = raw_text.split('\x0c') if '\x0c' in raw_text else [raw_text]
    unified_lines: List[str] = []

    for page in pages:
        lines = page.splitlines()
        banner_end = -1
        for idx, l in enumerate(lines[:15]):
            stripped = l.strip()
            if stripped and set(stripped) == {'-'}:
                banner_end = idx
                break

        if banner_end != -1:
            page_content = lines[banner_end + 1:]
        else:
            page_content = lines

        for line in page_content:
            if any(marker in line for marker in [
                'WAPDA ELECTRICITY BILLING',
                'MASTER DATA CHANGE',
                'POSTING LIST FOR MAIN DATE',
                'PROOF LIST FOR MAIN DATE',
                'DIVISON TOTAL:',
                'GRAND TOTAL:',
                'Total Records For the MainDate',
                'Total Records Posted',
                'Total Records UnPosted',
            ]):
                continue
            unified_lines.append(line)

    return unified_lines


def parse_single_txt_file(filepath: str) -> List[Dict[str, Any]]:
    """
    Parses a single CP-136 or CP-137 change posting text file and extracts structured records.
    Each record contains 'Ref No', 'Status', 'Report Type', and values for modified columns.
    """
    if not os.path.exists(filepath):
        raise FileNotFoundError(f"File not found: {filepath}")

    with open(filepath, 'r', encoding='utf-8', errors='ignore') as f:
        raw_text = f.read()

    base_filename = os.path.basename(filepath)
    if re.search(r'\b(?:RCP-137|CP-137|CP137)\b', raw_text[:2000], re.I) or '137' in base_filename:
        report_type = 'CP-137'
    else:
        report_type = 'CP-136'

    unified_lines = clean_page_banners(raw_text)
    ref_pattern = re.compile(r'^\s*(\d{14})')
    err_pattern = re.compile(r'\b([A-Z0-9]{2,5},)')

    ref_indices: List[int] = []
    for idx, line in enumerate(unified_lines):
        if ref_pattern.match(line):
            ref_indices.append(idx)

    records: List[Dict[str, Any]] = []

    ignored_b1 = ['CONN DATE', 'TOTAL', '----------------']
    ignored_b2 = [
        'METER RENT', 'CONTACT NO', 'CONSUMER NAME', 'CONNECTION CATEGORY',
        'FATHER/HUSBAND', 'ADDRES LINE', 'TRANSFORMER', 'NET MTR', 'SIM ID', 'SIM NO',
        'SET NO', '----------------'
    ]
    ignored_b3 = [
        'SET NO', 'CODE      NO', '----------------', 'New Ref No',
        'POSTED', 'NOT POSTED', 'CANNOT BE POSTED', 'READY FOR POSTING'
    ]

    spans_b3 = SPANS_BLOCK_3_CP136 if report_type == 'CP-136' else SPANS_BLOCK_3_CP137

    for idx, r_i in enumerate(ref_indices):
        next_i = ref_indices[idx + 1] if idx + 1 < len(ref_indices) else len(unified_lines)
        rec_lines = unified_lines[r_i:next_i]

        first_line = rec_lines[0]
        m_ref = ref_pattern.match(first_line)
        ref_no = m_ref.group(1) if m_ref else ''

        after_ref = first_line.lstrip()[14:].strip()
        customer_id = ''
        if after_ref:
            cust_match = re.match(r'^(\d+)', after_ref)
            if cust_match:
                customer_id = cust_match.group(1)

        # Status
        status = ''
        for line in rec_lines:
            line_u = line.upper()
            if 'CANNOT BE POSTED' in line_u or 'NOT POSTED' in line_u:
                status = 'NOT POSTED'
                break
            elif 'READY FOR POSTING' in line_u or ('POSTED' in line_u and 'TOTAL RECORDS' not in line_u):
                status = 'POSTED'

        # Error codes
        error_codes: List[str] = []
        for line in rec_lines:
            if any(k in line for k in ['Ref No', 'REF NO', 'SET NO', 'ERROR CODES', 'ERRORS CODES', 'PRINTING DATE', 'WAPDA']):
                continue
            if len(line) > 60:
                found = err_pattern.findall(line)
                for code in found:
                    code_clean = code.strip(',')
                    if code_clean not in error_codes and code_clean not in ['POSTED', 'TOTAL']:
                        error_codes.append(code_clean)

        record_data: Dict[str, Any] = {
            'Ref No': ref_no,
            'Sub Division': ref_no[2:7] if len(ref_no) >= 7 else '',
            'Status': status,
            'Report Type': report_type,
            'Customer ID': customer_id,
            'File Name': base_filename,
        }

        if error_codes:
            record_data['ERRORS CODES'] = ', '.join(error_codes)

        # Block 1 (CONN DATE ... GST)
        spans_b1 = SPANS_BLOCK_1_CP136 if report_type == 'CP-136' else SPANS_BLOCK_1_CP137
        for j, line in enumerate(rec_lines):
            if any(k in line for k in ['CONN DATE', 'DATE OF', 'DATE OF CONNECTION']):
                for k in range(j + 1, len(rec_lines)):
                    if 'METER RENT' in rec_lines[k] or '----------------' in rec_lines[k] or 'METER SERVICE' in rec_lines[k]:
                        break
                    s = rec_lines[k].strip()
                    if not s or any(ign in s for ign in ignored_b1):
                        continue
                    for col_name, start, end in spans_b1:
                        val = rec_lines[k][start:end].strip() if start < len(rec_lines[k]) else ''
                        if val:
                            record_data[col_name] = val
                break

        # Block 2 (METER RENT ... TRANSFRMR CODE)
        spans_b2 = SPANS_BLOCK_2_CP136 if report_type == 'CP-136' else SPANS_BLOCK_2_CP137
        for j, line in enumerate(rec_lines):
            if 'METER RENT' in line or 'METER SERVICE' in line:
                for k in range(j + 1, len(rec_lines)):
                    if 'SET NO' in rec_lines[k] or any(lbl in rec_lines[k] for lbl in ['CONSUMER NAME', 'CONTACT #', 'CONTACT NO', 'EMAIL ADDRESS', 'CONNECTION CATEGORY']):
                        break
                    s = rec_lines[k].strip()
                    if not s or any(ign in rec_lines[k] for ign in ignored_b2):
                        continue
                    for col_name, start, end in spans_b2:
                        val = rec_lines[k][start:end].strip() if start < len(rec_lines[k]) else ''
                        if val:
                            record_data[col_name] = val
                break

        # Key-Value labels in rec_lines (extracted against headings only)
        stop_lookahead = r'(?=\s*(?:EMAIL ADDRESS|CONNECTION CATEGORY|TRANSFORMER|NET MTR|LICENSE|SIM ID|SIM NO|\Z))'
        for line in rec_lines:
            if 'NET MTR CD' in line or 'NET MTR C' in line:
                m_cap = re.search(r'NET MTR\s+(?:CAP|C)\s*[:]?\s*([0-9.]+)', line)
                if m_cap:
                    record_data['NET MTR CAP'] = m_cap.group(1)
                m_cd = re.search(r'NET MTR\s+CD\s*[:]?\s*([0-9A-Z]+)', line)
                if m_cd:
                    record_data['NET MTR CD'] = m_cd.group(1)
            if 'NET MTR CONNECTION DAT' in line or 'NET MTR CONNECTION DATE' in line:
                m_conn = re.search(r'NET MTR CONNECTION DAT[E]?\s*[:]?\s*([0-9/:\sA-Za-z]+?)' + stop_lookahead, line)
                if m_conn and m_conn.group(1).strip():
                    record_data['NET MTR CONNECTION DATE'] = m_conn.group(1).strip()
            if 'NET MTR EXPIRY DA' in line or 'NET MTR EXPIRY DATE' in line:
                m_exp = re.search(r'NET MTR EXPIRY DA[TE]?\s*[:]?\s*([0-9/:\sA-Za-z]+?)' + stop_lookahead, line)
                if m_exp and m_exp.group(1).strip():
                    record_data['NET MTR EXPIRY DATE'] = m_exp.group(1).strip()
            elif 'EXPIRY DATE' in line and 'NET MTR' not in line:
                m_exp2 = re.search(r'EXPIRY\s+DATE\s*[:]?\s*([0-9/:\sA-Za-z]+?)' + stop_lookahead, line)
                if m_exp2 and m_exp2.group(1).strip():
                    record_data['EXPIRY DATE'] = m_exp2.group(1).strip()
            if 'LICENSE#' in line or 'LICENSE #' in line:
                m_lic = re.search(r'LICENSE\s*#?\s*[:]?\s*([A-Za-z0-9/_ -]+?)' + stop_lookahead, line)
                if m_lic and m_lic.group(1).strip():
                    record_data['LICENSE#'] = m_lic.group(1).strip()
            if 'CONNECTION CATEGORY' in line:
                m_cat = re.search(r'CONNECTION\s+CATEGORY\s+CODE\s*[:]?\s*([^:\n\r]+?)' + stop_lookahead, line)
                if m_cat and m_cat.group(1).strip():
                    record_data['CONNECTION CATEGORY CODE'] = m_cat.group(1).strip()
            if 'TRANSFORMER REF NO' in line:
                m_tr = re.search(r'TRANSFORMER\s+REF\s+NO\s*[:]?\s*([^:\n\r]+?)' + stop_lookahead, line)
                if m_tr and m_tr.group(1).strip():
                    record_data['TRANSFORMER REF NO'] = m_tr.group(1).strip()
            if 'TRANSFORMER ADDRESS' in line:
                m_ta = re.search(r'TRANSFORMER\s+ADDRESS\s*[:]?\s*([^:\n\r]+?)' + stop_lookahead, line)
                if m_ta and m_ta.group(1).strip():
                    record_data['TRANSFORMER ADDRESS'] = m_ta.group(1).strip()
            if 'CONSUMER NAME' in line:
                m_cn = re.search(r'CONSUMER\s+NAME\s*:\s*([^:\n\r]+?)' + stop_lookahead, line)
                if m_cn and m_cn.group(1).strip():
                    record_data['CONSUMER NAME'] = m_cn.group(1).strip()
            if 'FATHER/HUSBAND NAME' in line:
                m_fn = re.search(r'FATHER/HUSBAND\s+NAME\s*:\s*([^:\n\r]+?)' + stop_lookahead, line)
                if m_fn and m_fn.group(1).strip():
                    record_data['FATHER/HUSBAND NAME'] = m_fn.group(1).strip()
            if 'ADDRES LINE -1' in line or 'ADDRESS LINE -1' in line:
                m_a1 = re.search(r'ADDRES(?:S)?\s+LINE\s*-?1\s*:\s*([^:\n\r]+?)' + stop_lookahead, line)
                if m_a1 and m_a1.group(1).strip():
                    record_data['ADDRES LINE -1'] = m_a1.group(1).strip()
            if 'ADDRES LINE -2' in line or 'ADDRESS LINE -2' in line:
                m_a2 = re.search(r'ADDRES(?:S)?\s+LINE\s*-?2\s*:\s*([^:\n\r]+?)' + stop_lookahead, line)
                if m_a2 and m_a2.group(1).strip():
                    record_data['ADDRES LINE -2'] = m_a2.group(1).strip()
            if 'CONTACT #' in line or 'CONTACT NO' in line:
                m_ct = re.search(r'CONTACT\s*(?:#|NO)\s*:\s*([^:\n\r]+?)' + stop_lookahead, line)
                if m_ct and m_ct.group(1).strip():
                    record_data['CONTACT #'] = m_ct.group(1).strip()
            if 'EMAIL ADDRESS' in line:
                m_em = re.search(r'EMAIL\s+ADDRESS\s*:\s*([^:\n\r]+?)' + stop_lookahead, line)
                if m_em and m_em.group(1).strip():
                    record_data['EMAIL ADDRESS'] = m_em.group(1).strip()
            if 'SIM ID' in line:
                m_si = re.search(r'SIM\s+ID\s*:\s*([^:\n\r]+?)' + stop_lookahead, line)
                if m_si and m_si.group(1).strip():
                    record_data['SIM ID'] = m_si.group(1).strip()
            if 'SIM NO' in line:
                m_sn = re.search(r'SIM\s+NO\s*:\s*([^:\n\r]+?)' + stop_lookahead, line)
                if m_sn and m_sn.group(1).strip():
                    record_data['SIM NO'] = m_sn.group(1).strip()
            if 'NEW-REFERENCE-NUMBER' in line or 'New Ref No' in line:
                m_nr = re.search(r'(?:NEW-REFERENCE-NUMBER|New\s+Ref\s+No)\s*[:]?\s*(\d{14})', line)
                if m_nr:
                    record_data['NEW-REFERENCE-NUMBER'] = m_nr.group(1)

        # Block 3 (SET NO ... METER READING)
        set_lines: List[str] = []
        found_set = False
        for line in rec_lines:
            if 'SET NO' in line:
                found_set = True
                continue
            if found_set:
                if any(k in line for k in ['New Ref No', 'NEW-REFERENCE-NUMBER', 'POSTED', 'NOT POSTED', 'CANNOT BE POSTED', 'READY FOR POSTING', 'Ref No']):
                    break
                s = line.strip()
                if not s or any(ign in s for ign in ignored_b3):
                    continue
                set_lines.append(line)

        if set_lines:
            for sline in set_lines:
                row_copy = dict(record_data)
                for col_name, start, end in spans_b3:
                    if col_name == 'ERRORS CODES':
                        continue
                    val = sline[start:end].strip() if start < len(sline) else ''
                    if val:
                        row_copy[col_name] = val
                records.append(row_copy)
        else:
            records.append(record_data)

    return records


LABEL_DEFINITIONS: List[Tuple[str, List[str]]] = [
    ('NET MTR CONNECTION DATE', ['NET', 'MTR', 'CONNECTION', 'DATE']),
    ('NET MTR EXPIRY DATE', ['NET', 'MTR', 'EXPIRY', 'DATE']),
    ('CONNECTION CATEGORY CODE', ['CONNECTION', 'CATEGORY', 'CODE']),
    ('TRANSFORMER REF NO', ['TRANSFORMER', 'REF', 'NO']),
    ('TRANSFORMER ADDRESS', ['TRANSFORMER', 'ADDRESS']),
    ('ADDRES LINE -1', ['ADDRES', 'LINE', '-1']),
    ('ADDRES LINE -2', ['ADDRES', 'LINE', '-2']),
    ('ADDRESS LINE -1', ['ADDRESS', 'LINE', '-1']),
    ('ADDRESS LINE -2', ['ADDRESS', 'LINE', '-2']),
    ('ADDRES LINE -1', ['ADDRES', 'LINE', '1']),
    ('ADDRES LINE -2', ['ADDRES', 'LINE', '2']),
    ('NET MTR CD', ['NET', 'MTR', 'CD']),
    ('NET MTR CAP', ['NET', 'MTR', 'CAP']),
    ('CONSUMER NAME', ['CONSUMER', 'NAME']),
    ('FATHER/HUSBAND NAME', ['FATHER/HUSBAND', 'NAME']),
    ('CONTACT #', ['CONTACT', '#']),
    ('CONTACT NO', ['CONTACT', 'NO']),
    ('EMAIL ADDRESS', ['EMAIL', 'ADDRESS']),
    ('SIM ID', ['SIM', 'ID']),
    ('SIM NO', ['SIM', 'NO']),
    ('EXPIRY DATE', ['EXPIRY', 'DATE']),
    ('LICENSE#', ['LICENSE#']),
    ('LICENSE#', ['LICENSE', '#']),
]

HEADER_WORDS_METER = {
    'SET', 'NO', 'NO.', 'METER', 'TYPE', 'MAIN/', 'MAIN', 'S.', 'SET', 'KVARH/', 'KVAH-CD',
    'FEEDER', 'CODE', 'LOC', 'MFG', 'FACTOR', 'RANGE', 'RNG', 'READING', 'ERRORS', 'CODES', 'M/PLY', 'MTR'
}


def parse_single_pdf_file(filepath: str) -> List[Dict[str, Any]]:
    """
    Parses a single CP-136 or CP-137 change posting PDF report and extracts structured records.
    Accurately extracts 14-digit Ref Nos, Customer IDs, Status, Net Metering changes, Meter Sets,
    New Reference Numbers, Error Codes, and all modified fields without column contamination.
    """
    if not os.path.exists(filepath):
        raise FileNotFoundError(f"File not found: {filepath}")

    base_filename = os.path.basename(filepath)
    records: List[Dict[str, Any]] = []

    if pdfplumber is None:
        raise ImportError("pdfplumber is required for parsing CP-136 and CP-137 PDF reports.")

    with pdfplumber.open(filepath) as pdf:
        for pno, page in enumerate(pdf.pages):
            text = page.extract_text() or ''
            if not text or any(m in text for m in ['DIVISON TOTAL', 'GRAND TOTAL', 'Total Records For the MainDate', 'Total Records Posted']):
                continue

            words = page.extract_words()
            if not words:
                continue

            # Determine Report Type from page text (avoiding false positive on page numbers like PAGE: 137)
            banner_area = text[:500]
            if re.search(r'\b(?:RCP-137|CP-137|CP137)\b', banner_area, re.IGNORECASE):
                report_type = 'CP-137'
            elif re.search(r'\b(?:RCP-136|CP-136|CP136)\b', banner_area, re.IGNORECASE):
                report_type = 'CP-136'
            elif '137' in base_filename:
                report_type = 'CP-137'
            else:
                report_type = 'CP-136'

            # Determine Status from status text
            status = ''
            if re.search(r'CANNOT\s+BE\s+POSTED', text, re.IGNORECASE) or re.search(r'\bNOT\s+POSTED\b', text, re.IGNORECASE):
                status = 'NOT POSTED'
            elif re.search(r'READY\s+FOR\s+POSTING', text, re.IGNORECASE) or re.search(r'\bPOSTED\b', text, re.IGNORECASE):
                status = 'POSTED'

            # Group words by Y line (rounded to 3pt)
            lines: Dict[float, List[Dict[str, Any]]] = {}
            for w in words:
                top_k = round(w['top'] / 3.0) * 3.0
                lines.setdefault(top_k, []).append(w)

            # Landmarks
            y_top_hdr = None
            y_b1_hdr = None
            y_b2_hdr = None
            y_mid_hdr = None
            y_meter_hdr = None
            y_newref_hdr = None

            for top_k in sorted(lines.keys()):
                line_words = sorted(lines[top_k], key=lambda x: x['x0'])
                line_text = ' '.join(w['text'].upper() for w in line_words)
                
                if ('REF NO' in line_text or 'REF. NO' in line_text) and top_k < 140 and y_top_hdr is None:
                    y_top_hdr = top_k
                elif any(k in line_text for k in ['CONN DATE', 'DATE OF', 'DATE OF CONNECTION']) and 120 < top_k < 180 and y_b1_hdr is None:
                    y_b1_hdr = top_k
                elif ('METER RENT' in line_text or ('METER' in line_text and 'SERVICE' in line_text)) and 170 < top_k < 235 and y_b2_hdr is None:
                    y_b2_hdr = top_k
                elif any(k in line_text for k in ['CONSUMER NAME', 'CONTACT #', 'CONTACT NO', 'EMAIL ADDRESS']) and 200 < top_k < 290 and y_mid_hdr is None:
                    y_mid_hdr = top_k
                elif 'SET NO' in line_text and 290 < top_k < 430 and y_meter_hdr is None:
                    y_meter_hdr = top_k
                elif any(k in line_text for k in ['NEW REF NO', 'NEW-REFERENCE-NUMBER']) and top_k > 320 and y_newref_hdr is None:
                    y_newref_hdr = top_k

            # Find Ref No (14 digits)
            ref_no = ''
            y_ref = None
            for top_k in sorted(lines.keys()):
                if y_top_hdr and top_k < y_top_hdr:
                    continue
                if y_b1_hdr and top_k >= y_b1_hdr:
                    break
                for w in lines[top_k]:
                    if re.match(r'^\d{14}$', w['text']) and w['x0'] < 120:
                        ref_no = w['text']
                        y_ref = top_k
                        break
                if ref_no:
                    break

            if not ref_no:
                continue

            all_error_codes: List[str] = []

            rec: Dict[str, Any] = {
                'Ref No': ref_no,
                'Sub Division': ref_no[2:7] if len(ref_no) >= 7 else '',
                'Status': status,
                'Report Type': report_type,
                'File Name': base_filename,
            }

            # Top Block data words (below y_top_hdr, strictly before y_b1_hdr)
            top_data_words = []
            y_b1_limit = y_b1_hdr if y_b1_hdr else 160
            for yk in lines:
                if (y_top_hdr and yk > y_top_hdr + 6) and yk < y_b1_limit:
                    for w in lines[yk]:
                        if w['text'] != ref_no and w['text'].upper() not in ['(', 'OLD', ')', '(OLD)', 'TOTAL', 'GRAND', 'SUB', 'DIVISION']:
                            top_data_words.append(w)
            top_data_words = sorted(top_data_words, key=lambda x: x['x0'])

            for w in top_data_words:
                txt = w['text']
                x = w['x0']
                if re.match(r'^[A-Z0-9]+,', txt) or (x >= 630 and not txt.isdigit()):
                    for ec in txt.split(','):
                        ec = ec.strip()
                        if ec and ec not in all_error_codes and ec not in ['POSTED', 'NOT']:
                            all_error_codes.append(ec)
                elif txt.isdigit() and len(txt) >= 8 and (x < 150 or (report_type == 'CP-136' and 160 <= x < 240)):
                    rec['Customer ID'] = txt
                elif report_type == 'CP-136':
                    # CP-136 top columns: REF NO, SUB BATCH, CUSTOMER ID, TARIFF, APPLICATION NO, APPL DATE, SEC AMOUNT, SEC DATE, ERROR CODES
                    if 100 <= x < 165: rec['SUB BATCH'] = txt
                    elif 235 <= x < 285: rec['TARIFF'] = txt
                    elif 285 <= x < 370: rec['APPLICATION NO'] = txt
                    elif 370 <= x < 445: rec['APP DATE'] = txt
                    elif 445 <= x < 530: rec['SEC AMOUNT'] = txt
                    elif 530 <= x < 620: rec['SEC DATE'] = txt
                else:  # CP-137
                    # CP-137 top columns: Ref. No, Customer ID, SUB BATCH, TARIFF, APP. NO, APP. DATE, SEC. AMNT, SEC. DATE, MAIN SETS, SUB SETS, S-POINTS, ERROR CODES
                    if 145 <= x < 205: rec['SUB BATCH'] = txt
                    elif 205 <= x < 245: rec['TARIFF'] = txt
                    elif 245 <= x < 300: rec['APPLICATION NO'] = txt
                    elif 300 <= x < 365: rec['APP DATE'] = txt
                    elif 365 <= x < 440: rec['SEC AMOUNT'] = txt
                    elif 440 <= x < 485: rec['SEC DATE'] = txt
                    elif 485 <= x < 545 and txt.isdigit(): rec['MAIN SETS'] = txt
                    elif 545 <= x < 600 and txt.isdigit(): rec['SUB SETS'] = txt
                    elif 600 <= x < 645 and not re.match(r'^[A-Z0-9]+,', txt): rec['S-POINTS'] = txt

            # Block 1 (below headings, strictly before y_b2_hdr)
            HEADER_WORDS_B1 = {
                'DATE', 'OF', 'CONNECTION', 'SANC', 'SAN', 'LOAD', 'TRANS.', 'TRANS', 'CAPCTY',
                'R/U', 'RU', 'F/P', 'SEASON', 'CD', 'CONS.', 'CONS', 'CLASS', 'ELEC.', 'ELEC',
                'DUTY', 'GOVT.', 'GOVT', 'DEPT.', 'DEPT', 'GVT', 'DEP', 'STD-CLSF.', 'STD-CLSF', 'STD',
                'CLSF', 'I-TAX', 'IT', 'TOTAL', 'METERS', 'AC', 'NTN', 'GST', 'OLD', '(', ')'
            }
            if y_b1_hdr is not None:
                hdr_offset = 18 if report_type == 'CP-137' else 14
                y_max_b1 = (y_b2_hdr - 6) if y_b2_hdr else (y_b1_hdr + 40)
                b1_words = [
                    w for w in words
                    if y_b1_hdr + hdr_offset <= w['top'] < y_max_b1
                    and not w['text'].startswith('---')
                    and w['text'].upper().rstrip('.') not in HEADER_WORDS_B1
                ]
                for w in b1_words:
                    x = w['x0']
                    txt = w['text']
                    if re.match(r'^[A-Z0-9]+,', txt) or (x >= 630 and x >= 600 and re.match(r'^[A-Z0-9]{2,5},', txt)):
                        for ec in txt.split(','):
                            ec = ec.strip()
                            if ec and ec not in all_error_codes and ec not in ['POSTED', 'NOT']:
                                all_error_codes.append(ec)
                    elif report_type == 'CP-136':
                        if 20 <= x < 85: rec['CONN DATE'] = txt
                        elif 85 <= x < 155: rec['SAN LOAD'] = txt
                        elif 155 <= x < 200: rec['RU'] = txt
                        elif 200 <= x < 240: rec['F/P'] = txt
                        elif 240 <= x < 335: rec['SEASON CD'] = txt
                        elif 335 <= x < 385: rec['ED CD'] = txt
                        elif 385 <= x < 445: rec['GVT DEP'] = txt
                        elif 445 <= x < 500: rec['STD CLSF'] = txt
                        elif 500 <= x < 535: rec['IT'] = txt
                        elif 535 <= x < 595: rec['METERS'] = txt
                        elif 595 <= x < 630: rec['AC'] = txt
                        elif 630 <= x < 685: rec['NTN'] = txt
                        elif 685 <= x < 750: rec['GST'] = txt
                    else:  # CP-137
                        if 20 <= x < 115: rec['CONN DATE'] = txt
                        elif 115 <= x < 160: rec['SAN LOAD'] = txt
                        elif 160 <= x < 200: rec['TRANS. CAPCTY'] = txt
                        elif 200 <= x < 225: rec['RU'] = txt
                        elif 225 <= x < 250: rec['F/P'] = txt
                        elif 250 <= x < 300: rec['SEASON CD'] = txt
                        elif 300 <= x < 345: rec['CONS. CLASS'] = txt
                        elif 345 <= x < 380: rec['ED CD'] = txt
                        elif 380 <= x < 425: rec['GVT DEP'] = txt
                        elif 425 <= x < 485: rec['STD CLSF'] = txt
                        elif 485 <= x < 525: rec['IT'] = txt
                        elif 525 <= x < 600: rec['METERS'] = txt
                        elif 600 <= x < 630: rec['AC'] = txt
                        elif 630 <= x < 680: rec['NTN'] = txt
                        elif 680 <= x < 750: rec['GST'] = txt

            # Block 2 (below headings, strictly before middle label section)
            HEADER_WORDS_B2 = {
                'METER', 'RENT', 'SERVICE', 'SRV', 'N.', 'N', 'ID.', 'ID', 'NUMBER',
                'TRANSFRMR.', 'TRANSFRMR', 'TRANSFORMER', 'CODE', 'STREET', 'LIGHT', 'CATEGORY',
                'LINE', 'LENGTH', 'OLD', '(', ')', 'CONTACT', 'NO', 'EMAIL', 'ADDRESS', 'N.I.D',
                '<-----------STREET', 'LIGHT----------->'
            }
            if y_b2_hdr is not None:
                for yk in lines:
                    if abs(yk - y_b2_hdr) <= 6:
                        for w in lines[yk]:
                            if w['x0'] >= 600 and re.match(r'^[A-Z0-9]+,', w['text']):
                                for ec in w['text'].split(','):
                                    ec = ec.strip()
                                    if ec and ec not in all_error_codes and ec not in ['POSTED', 'NOT']:
                                        all_error_codes.append(ec)

                hdr_offset = 18 if report_type == 'CP-137' else 14
                y_max_b2 = (y_mid_hdr - 6) if y_mid_hdr else (y_b2_hdr + 35)
                b2_words = [
                    w for w in words
                    if y_b2_hdr + hdr_offset <= w['top'] < y_max_b2
                    and not w['text'].startswith('---')
                    and w['text'].upper().rstrip('.') not in HEADER_WORDS_B2
                ]
                for w in b2_words:
                    x = w['x0']
                    txt = w['text']
                    if re.match(r'^[A-Z0-9]+,', txt):
                        for ec in txt.split(','):
                            ec = ec.strip()
                            if ec and ec not in all_error_codes and ec not in ['POSTED', 'NOT']:
                                all_error_codes.append(ec)
                    elif report_type == 'CP-136':
                        if 20 <= x < 100: rec['METER RENT'] = txt
                        elif 100 <= x < 170: rec['SRV RENT'] = txt
                        elif 170 <= x < 260: rec['N .ID NUMBER'] = txt
                        elif 260 <= x < 370: rec['TRANSFRMR CODE'] = txt
                    else:  # CP-137
                        if 20 <= x < 140: rec['METER RENT'] = txt
                        elif 140 <= x < 210: rec['SRV RENT'] = txt
                        elif 210 <= x < 285: rec['N .ID NUMBER'] = txt
                        elif 285 <= x < 380: rec['TRANSFRMR CODE'] = txt
                        elif 440 <= x < 510: rec['STREET LIGHT CATEGORY'] = txt
                        elif 510 <= x < 600: rec['STREET LIGHT LINE LENGTH'] = txt

            # Middle Label Section (AGAINST HEADINGS ONLY!)
            mid_words = []
            y_min_mid = (y_b2_hdr + 15) if y_b2_hdr else 210
            y_max_mid = (y_meter_hdr - 10) if y_meter_hdr else 395
            for yk in lines:
                if y_min_mid <= yk <= y_max_mid:
                    mid_words.extend(lines[yk])

            sorted_defs = sorted(LABEL_DEFINITIONS, key=lambda x: len(x[1]), reverse=True)
            found_labels = []
            claimed_word_indices = set()
            w_sorted = sorted(mid_words, key=lambda w: (round(w['top'] / 3.0), w['x0']))
            
            for label_name, token_list in sorted_defs:
                for i in range(len(w_sorted) - len(token_list) + 1):
                    if any(idx in claimed_word_indices for idx in range(i, i + len(token_list))):
                        continue
                    match = True
                    for k, tok in enumerate(token_list):
                        if w_sorted[i+k]['text'].upper().rstrip(':') != tok.upper().rstrip(':'):
                            match = False
                            break
                        if k > 0 and abs(w_sorted[i+k]['top'] - w_sorted[i]['top']) > 5:
                            match = False
                            break
                    if match:
                        first_w = w_sorted[i]
                        last_w = w_sorted[i + len(token_list) - 1]
                        found_labels.append({
                            'name': label_name,
                            'top': first_w['top'],
                            'x0': first_w['x0'],
                            'x1': last_w['x1']
                        })
                        for idx in range(i, i + len(token_list)):
                            claimed_word_indices.add(idx)

            for lbl in found_labels:
                y = lbl['top']
                x1 = lbl['x1']
                next_labels = [ol for ol in found_labels if ol != lbl and abs(ol['top'] - y) <= 5 and ol['x0'] > x1]
                x_limit = min([ol['x0'] for ol in next_labels]) if next_labels else 760

                val_words = []
                for w in mid_words:
                    if abs(w['top'] - y) <= 5 and (x1 + 1) < w['x0'] < x_limit:
                        if any(l['x0'] - 2 <= w['x0'] <= l['x1'] + 2 and abs(l['top'] - w['top']) <= 5 for l in found_labels):
                            continue
                        if re.match(r'^[A-Z0-9]+,', w['text']):
                            for ec in w['text'].split(','):
                                ec = ec.strip()
                                if ec and ec not in all_error_codes and ec not in ['POSTED', 'NOT']:
                                    all_error_codes.append(ec)
                            continue
                        val_words.append(w)

                if val_words:
                    val_str = ' '.join(w['text'] for w in sorted(val_words, key=lambda x: x['x0'])).strip().lstrip(':').strip()
                    if val_str and val_str not in rec:
                        norm_key = lbl['name']
                        if norm_key == 'ADDRESS LINE -1': norm_key = 'ADDRES LINE -1'
                        if norm_key == 'ADDRESS LINE -2': norm_key = 'ADDRES LINE -2'
                        if norm_key == 'CONTACT NO': norm_key = 'CONTACT #'
                        rec[norm_key] = val_str

            # Meter Table rows (below y_meter_hdr, strictly before y_newref_hdr)
            meter_rows = []
            if y_meter_hdr is not None:
                y_max_m = (y_newref_hdr - 6) if y_newref_hdr else (y_meter_hdr + 80)
                m_data_lines = {}
                for yk in lines:
                    if y_meter_hdr + 18 <= yk < y_max_m:
                        lw = lines[yk]
                        line_str = ' '.join(w['text'].upper() for w in lw)
                        if any(k in line_str for k in ['NEW REF', 'CANNOT BE POSTED', 'READY FOR POSTING', 'NEW-REFERENCE']):
                            continue
                        if all(w['text'].upper().rstrip('.') in HEADER_WORDS_METER for w in lw):
                            continue
                        m_data_lines[yk] = sorted(lw, key=lambda x: x['x0'])

                for yk in sorted(m_data_lines.keys()):
                    lw = m_data_lines[yk]
                    mr = {}
                    for w in lw:
                        x = w['x0']
                        txt = w['text']
                        if re.match(r'^[A-Z0-9]+,', txt) or (x >= 640 and report_type == 'CP-137' and not txt.isdigit()) or (x >= 540 and report_type == 'CP-136' and not txt.isdigit()):
                            for ec in txt.split(','):
                                ec = ec.strip()
                                if ec and ec not in all_error_codes and ec not in ['POSTED', 'NOT']:
                                    all_error_codes.append(ec)
                        elif txt.upper() in ['POSTED', 'NOT', 'CANNOT', 'BE', 'READY', 'FOR']:
                            continue
                        elif report_type == 'CP-136':
                            if 15 <= x < 55 and txt.isdigit(): mr['SET NO'] = txt
                            elif 55 <= x < 110: mr['METER TYPE'] = txt
                            elif 110 <= x < 185: mr['FEEDER CODE'] = txt
                            elif 185 <= x < 245: mr['METER NO'] = txt
                            elif 245 <= x < 290: mr['LOC CODE'] = txt
                            elif 290 <= x < 340: mr['MFG CODE'] = txt
                            elif 340 <= x < 390: mr['METER CODE'] = txt
                            elif 390 <= x < 450: mr['M/PLY FACTOR'] = txt
                            elif 450 <= x < 485: mr['MTR RANGE'] = txt
                            elif 485 <= x < 545: mr['METER READING'] = txt
                        else:  # CP-137
                            if 15 <= x < 55 and txt.isdigit(): mr['SET NO'] = txt
                            elif 55 <= x < 100: mr['METER TYPE'] = txt
                            elif 100 <= x < 150: mr['MAIN/S. SET'] = txt
                            elif 150 <= x < 205: mr['KVARH/KVAH-CD'] = txt
                            elif 205 <= x < 265: mr['FEEDER CODE'] = txt
                            elif 265 <= x < 335: mr['METER NO'] = txt
                            elif 335 <= x < 375: mr['LOC CODE'] = txt
                            elif 375 <= x < 415: mr['MFG CODE'] = txt
                            elif 415 <= x < 465: mr['METER CODE'] = txt
                            elif 465 <= x < 510: mr['M/PLY FACTOR'] = txt
                            elif 510 <= x < 555: mr['MTR RANGE'] = txt
                            elif 555 <= x < 640: mr['METER READING'] = txt

                    if mr:
                        meter_rows.append(mr)

            # NEW-REFERENCE-NUMBER block
            if y_newref_hdr:
                nr_words = [
                    w for w in words
                    if y_newref_hdr + 4 <= w['top'] < y_newref_hdr + 45
                    and not w['text'].startswith('---')
                    and w['text'].upper() not in ['CANNOT', 'BE', 'POSTED', 'READY', 'FOR', 'POSTING']
                ]
                for w in nr_words:
                    x = w['x0']
                    txt = w['text']
                    if x < 160 and re.match(r'^\d{14}$', txt):
                        rec['NEW-REFERENCE-NUMBER'] = txt
                    elif 160 <= x < 235:
                        rec['A/CODE'] = txt
                    elif 235 <= x < 480:
                        rec['BANK NAME'] = (rec.get('BANK NAME', '') + ' ' + txt).strip()
                    elif 480 <= x < 565:
                        rec['AMOUNT'] = txt
                    elif 565 <= x < 640:
                        rec['START DATE'] = txt
                    elif 640 <= x < 740:
                        rec['EXPIRY DATE'] = txt

            # Attach consolidated error codes
            if all_error_codes:
                rec['ERRORS CODES'] = ', '.join(all_error_codes)

            if meter_rows:
                for mr in meter_rows:
                    rc = dict(rec)
                    rc.update(mr)
                    records.append(rc)
            else:
                records.append(rec)

    return records


def parse_single_file(filepath: str) -> List[Dict[str, Any]]:
    """
    Parses a single CP-136 or CP-137 change posting report (.txt or .pdf) and extracts structured records.
    Each record contains 'Ref No', 'Status', 'Report Type', and values for modified columns.
    """
    if not os.path.exists(filepath):
        raise FileNotFoundError(f"File not found: {filepath}")

    if filepath.lower().endswith('.pdf'):
        return parse_single_pdf_file(filepath)
    return parse_single_txt_file(filepath)


def deduplicate_records(
    records: List[Dict[str, Any]],
    mode: str = 'prefer_posted'
) -> Tuple[List[Dict[str, Any]], int]:
    """
    Deduplicates records by Report Type, Reference Number ('Ref No'), and Meter Set ('SET NO').
    Ensures that only single postings remain in the output, deleting duplicate occurrences.
    """
    if not records:
        return [], 0

    seen: Dict[Tuple[str, str, str], Dict[str, Any]] = {}
    order: List[Tuple[str, str, str]] = []
    duplicates_removed = 0

    for rec in records:
        ref = rec.get('Ref No', '')
        if not ref:
            continue

        rtype = rec.get('Report Type', 'CP-136')
        set_no = rec.get('SET NO', '')
        key = (rtype, ref, set_no)

        if key not in seen:
            seen[key] = dict(rec)
            order.append(key)
        else:
            duplicates_removed += 1
            existing = seen[key]

            if mode == 'first':
                continue
            elif mode == 'last':
                seen[key] = dict(rec)
            else:  # 'prefer_posted' (default)
                if rec.get('Status') == 'POSTED' and existing.get('Status') != 'POSTED':
                    existing['Status'] = 'POSTED'
                    existing['Customer ID'] = rec.get('Customer ID', existing.get('Customer ID'))
                    existing['ERRORS CODES'] = rec.get('ERRORS CODES', '')
                    existing['File Name'] = rec.get('File Name', existing.get('File Name'))

                for k, v in rec.items():
                    if v and not existing.get(k):
                        existing[k] = v

    unique_records = [seen[k] for k in order]
    return unique_records, duplicates_removed


def parse_multiple_files(
    filepaths: List[str],
    deduplicate: bool = True,
    dedup_mode: str = 'prefer_posted'
) -> Tuple[List[Dict[str, Any]], Dict[str, Any]]:
    """
    Parses multiple CP-136 and CP-137 files and aggregates their records.
    Optionally deduplicates to keep single postings.

    Returns:
        all_records: list of parsed (and optionally deduplicated) records
        summary_stats: summary counts
    """
    raw_records: List[Dict[str, Any]] = []
    stats = {
        'total_files': len(filepaths),
        'raw_records': 0,
        'duplicates_removed': 0,
        'total_records': 0,
        'posted_records': 0,
        'not_posted_records': 0,
        'cp136_count': 0,
        'cp137_count': 0,
        'files_summary': {}
    }

    for path in filepaths:
        fname = os.path.basename(path)
        try:
            records = parse_single_file(path)
            posted = sum(1 for r in records if r.get('Status') == 'POSTED')
            not_posted = sum(1 for r in records if r.get('Status') == 'NOT POSTED')
            rtype = records[0].get('Report Type', 'CP-136') if records else ('CP-137' if '137' in fname else 'CP-136')

            raw_records.extend(records)
            stats['files_summary'][fname] = {
                'report_type': rtype,
                'records': len(records),
                'posted': posted,
                'not_posted': not_posted,
                'status': 'OK'
            }
        except Exception as e:
            stats['files_summary'][fname] = {
                'report_type': 'Unknown',
                'records': 0,
                'posted': 0,
                'not_posted': 0,
                'status': f'Error: {str(e)}'
            }

    stats['raw_records'] = len(raw_records)

    if deduplicate:
        final_records, dups_removed = deduplicate_records(raw_records, mode=dedup_mode)
        stats['duplicates_removed'] = dups_removed
    else:
        final_records = raw_records
        stats['duplicates_removed'] = 0

    stats['total_records'] = len(final_records)
    stats['posted_records'] = sum(1 for r in final_records if r.get('Status') == 'POSTED')
    stats['not_posted_records'] = sum(1 for r in final_records if r.get('Status') == 'NOT POSTED')
    stats['cp136_count'] = sum(1 for r in final_records if r.get('Report Type') == 'CP-136')
    stats['cp137_count'] = sum(1 for r in final_records if r.get('Report Type') == 'CP-137')

    return final_records, stats


def _write_data_sheet(
    ws,
    sheet_title: str,
    records: List[Dict[str, Any]],
    only_present_columns: bool = True,
    include_status: bool = True,
    include_filename: bool = False
):
    """
    Helper to populate and format a data worksheet.
    """
    ws.title = sheet_title
    ws.views.sheetView[0].showGridLines = True

    # Build active columns
    columns = ['Ref No', 'Sub Division']
    if include_status:
        columns.append('Status')
    if include_filename:
        columns.append('File Name')

    if only_present_columns:
        active_headings = [h for h in ALL_ORDERED_HEADINGS if h not in columns and any(rec.get(h) for rec in records)]
        # Add any extra keys
        for rec in records:
            for k in rec.keys():
                if k not in columns and k not in ALL_ORDERED_HEADINGS and k not in ['Report Type', 'Page'] and k not in active_headings:
                    active_headings.append(k)
    else:
        active_headings = [h for h in ALL_ORDERED_HEADINGS if h not in columns]

    columns.extend(active_headings)

    # Styles
    font_header = Font(name='Calibri', size=11, bold=True, color='FFFFFF')
    fill_header = PatternFill(start_color='1F4E78', end_color='1F4E78', fill_type='solid')

    font_data = Font(name='Calibri', size=10)
    fill_even = PatternFill(start_color='FFFFFF', end_color='FFFFFF', fill_type='solid')
    fill_odd = PatternFill(start_color='F4F7FA', end_color='F4F7FA', fill_type='solid')

    fill_posted = PatternFill(start_color='D4EDDA', end_color='D4EDDA', fill_type='solid')
    font_posted = Font(name='Calibri', size=10, bold=True, color='155724')

    fill_not_posted = PatternFill(start_color='F8D7DA', end_color='F8D7DA', fill_type='solid')
    font_not_posted = Font(name='Calibri', size=10, bold=True, color='721C24')

    align_center = Alignment(horizontal='center', vertical='center')
    align_left = Alignment(horizontal='left', vertical='center')
    align_header = Alignment(horizontal='center', vertical='center', wrap_text=True)

    border_thin = Border(
        left=Side(style='thin', color='D9D9D9'),
        right=Side(style='thin', color='D9D9D9'),
        top=Side(style='thin', color='D9D9D9'),
        bottom=Side(style='thin', color='D9D9D9')
    )
    border_header = Border(
        left=Side(style='thin', color='13324E'),
        right=Side(style='thin', color='13324E'),
        top=Side(style='thin', color='13324E'),
        bottom=Side(style='medium', color='0F2438')
    )

    # Header Row
    ws.row_dimensions[1].height = 28
    for col_idx, col_name in enumerate(columns, start=1):
        cell = ws.cell(row=1, column=col_idx, value=col_name)
        cell.font = font_header
        cell.fill = fill_header
        cell.alignment = align_header
        cell.border = border_header

    # Data Rows
    text_columns = {
        'Ref No', 'Customer ID', 'NEW-REFERENCE-NUMBER', 'METER NO',
        'N .ID NUMBER', 'SIM ID', 'SIM NO', 'LOC CODE', 'MFG CODE',
        'METER CODE', 'FEEDER CODE', 'TRANSFRMR CODE'
    }

    ref_col_idx = columns.index('Ref No') + 1 if 'Ref No' in columns else 1
    ref_letter = get_column_letter(ref_col_idx)

    for row_idx, rec in enumerate(records, start=2):
        ws.row_dimensions[row_idx].height = 20
        row_fill = fill_odd if (row_idx % 2 == 1) else fill_even

        for col_idx, col_name in enumerate(columns, start=1):
            val = rec.get(col_name, '')
            cell = ws.cell(row=row_idx, column=col_idx)

            if col_name == 'Sub Division':
                cell.value = f"=MID({ref_letter}{row_idx}, 3, 5)"
                cell.alignment = align_center
                cell.font = font_data
                cell.fill = row_fill
            elif col_name in text_columns:
                cell.number_format = '@'
                cell.value = str(val) if val else ''
                cell.alignment = align_center
                cell.font = font_data
                cell.fill = row_fill
            elif col_name == 'Status':
                cell.value = str(val) if val else ''
                cell.alignment = align_center
                if val == 'POSTED':
                    cell.fill = fill_posted
                    cell.font = font_posted
                elif val == 'NOT POSTED':
                    cell.fill = fill_not_posted
                    cell.font = font_not_posted
                else:
                    cell.fill = row_fill
                    cell.font = font_data
            else:
                cell.value = val
                cell.alignment = align_center if len(str(val)) <= 15 else align_left
                cell.font = font_data
                cell.fill = row_fill

            cell.border = border_thin

    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = ws.dimensions

    # Auto-adjust column widths
    for col_idx, col_name in enumerate(columns, start=1):
        col_letter = get_column_letter(col_idx)
        max_len = len(str(col_name))
        for row in range(2, min(len(records) + 2, 200)):
            cell_val = str(ws.cell(row=row, column=col_idx).value or '')
            if len(cell_val) > max_len:
                max_len = len(cell_val)
        ws.column_dimensions[col_letter].width = max(max_len + 4, 13)


def export_to_excel(
    records: List[Dict[str, Any]],
    output_path: str,
    only_present_columns: bool = True,
    include_status: bool = True,
    include_filename: bool = False,
    include_summary_sheet: bool = True,
    summary_stats: Dict[str, Any] = None
) -> str:
    """
    Exports parsed CP136 and CP137 records to an Excel workbook.
    If CP-137 records are present, they are placed in a dedicated 'CP137 Posting Data' worksheet.
    If both CP-136 and CP-137 are present, both sheets are generated in the same workbook.
    Formats Reference Numbers as text ('@') to prevent scientific notation.
    """
    wb = openpyxl.Workbook()

    cp136_records = [r for r in records if r.get('Report Type') == 'CP-136']
    cp137_records = [r for r in records if r.get('Report Type') == 'CP-137']

    # Fallback if report types weren't assigned
    if not cp136_records and not cp137_records:
        cp136_records = records

    first_sheet = True

    if cp136_records:
        ws_136 = wb.active if first_sheet else wb.create_sheet()
        _write_data_sheet(
            ws=ws_136,
            sheet_title='CP136 Posting Data',
            records=cp136_records,
            only_present_columns=only_present_columns,
            include_status=include_status,
            include_filename=include_filename
        )
        first_sheet = False

    if cp137_records:
        ws_137 = wb.active if first_sheet else wb.create_sheet()
        _write_data_sheet(
            ws=ws_137,
            sheet_title='CP137 Posting Data',
            records=cp137_records,
            only_present_columns=only_present_columns,
            include_status=include_status,
            include_filename=include_filename
        )
        first_sheet = False

    # Summary Sheet
    if include_summary_sheet and summary_stats:
        ws_sum = wb.create_sheet(title='Summary', index=0)
        ws_sum.views.sheetView[0].showGridLines = True

        title_font = Font(name='Calibri', size=14, bold=True, color='1F4E78')
        sec_font = Font(name='Calibri', size=11, bold=True)
        lbl_font = Font(name='Calibri', size=11, bold=True, color='333333')
        val_font = Font(name='Calibri', size=11)
        font_header = Font(name='Calibri', size=11, bold=True, color='FFFFFF')
        fill_header = PatternFill(start_color='1F4E78', end_color='1F4E78', fill_type='solid')
        align_center = Alignment(horizontal='center', vertical='center')
        align_left = Alignment(horizontal='left', vertical='center')

        ws_sum['A1'] = "CP136 & 137 Change Posting Report Summary"
        ws_sum['A1'].font = title_font
        ws_sum.row_dimensions[1].height = 28

        stats_items = [
            ("Total Files Processed", summary_stats.get('total_files', 0)),
        ]
        if summary_stats.get('duplicates_removed', 0) > 0:
            stats_items.extend([
                ("Total Records Found (Raw)", summary_stats.get('raw_records', summary_stats.get('total_records', 0))),
                ("Duplicate Records Removed", summary_stats.get('duplicates_removed', 0)),
                ("Unique Single Records Exported", summary_stats.get('total_records', len(records))),
            ])
        else:
            stats_items.append(("Total Records Extracted", summary_stats.get('total_records', len(records))))

        stats_items.extend([
            ("Total Records Posted", summary_stats.get('posted_records', 0)),
            ("Total Records Unposted (Not Posted)", summary_stats.get('not_posted_records', 0)),
        ])

        if cp136_records and cp137_records:
            stats_items.extend([
                ("CP-136 Records Exported", len(cp136_records)),
                ("CP-137 Records Exported", len(cp137_records)),
            ])

        for s_idx, (label, value) in enumerate(stats_items, start=3):
            c_lbl = ws_sum.cell(row=s_idx, column=1, value=label)
            c_lbl.font = lbl_font
            c_val = ws_sum.cell(row=s_idx, column=2, value=value)
            c_val.font = val_font
            c_val.alignment = align_center

        f_start = len(stats_items) + 5
        ws_sum.cell(row=f_start, column=1, value="Files Breakdown").font = sec_font

        f_hdr_row = f_start + 1
        ws_sum.row_dimensions[f_hdr_row].height = 24
        file_hdrs = ["Report Type", "File Name", "Records", "Posted", "Not Posted", "Status"]
        for f_c, h in enumerate(file_hdrs, start=1):
            cell = ws_sum.cell(row=f_hdr_row, column=f_c, value=h)
            cell.font = font_header
            cell.fill = fill_header
            cell.alignment = align_center

        f_data_row = f_hdr_row + 1
        for fname, f_info in summary_stats.get('files_summary', {}).items():
            ws_sum.cell(row=f_data_row, column=1, value=f_info.get('report_type', 'CP-136')).alignment = align_center
            ws_sum.cell(row=f_data_row, column=2, value=fname).alignment = align_left
            ws_sum.cell(row=f_data_row, column=3, value=f_info.get('records', 0)).alignment = align_center
            ws_sum.cell(row=f_data_row, column=4, value=f_info.get('posted', 0)).alignment = align_center
            ws_sum.cell(row=f_data_row, column=5, value=f_info.get('not_posted', 0)).alignment = align_center
            ws_sum.cell(row=f_data_row, column=6, value=f_info.get('status', 'OK')).alignment = align_center
            f_data_row += 1

        ws_sum.column_dimensions['A'].width = 18
        ws_sum.column_dimensions['B'].width = 38
        ws_sum.column_dimensions['C'].width = 15
        ws_sum.column_dimensions['D'].width = 15
        ws_sum.column_dimensions['E'].width = 15
        ws_sum.column_dimensions['F'].width = 15

    # Ensure output directory exists
    out_dir = os.path.dirname(os.path.abspath(output_path))
    if out_dir and not os.path.exists(out_dir):
        os.makedirs(out_dir, exist_ok=True)

    wb.save(output_path)
    return os.path.abspath(output_path)


def run_cp136_check(
    file_paths: List[str],
    output_path: Optional[str] = None,
    deduplicate: bool = True,
    dedup_mode: str = 'prefer_posted',
    only_present: bool = True,
    include_status: bool = True,
    include_filename: bool = False,
    log_fn: Optional[Callable[[str], None]] = None
) -> Tuple[str, Dict[str, Any]]:
    """
    High-level runner used by PITC Suite Worker.
    Returns (saved_excel_path, stats_dict).
    """
    def _log(msg: str):
        if log_fn:
            log_fn(msg)

    if not file_paths:
        raise ValueError("Please select at least one CP-136 or CP-137 report file (.txt or .pdf).")

    _log(f"Processing {len(file_paths)} report file(s)...")

    records, stats = parse_multiple_files(
        filepaths=file_paths,
        deduplicate=deduplicate,
        dedup_mode=dedup_mode
    )

    if not records:
        raise ValueError("No records found in the selected file(s).")

    _log(f"Found {stats.get('raw_records', len(records))} raw record(s).")
    if stats.get('duplicates_removed', 0) > 0:
        _log(f"Removed {stats['duplicates_removed']} duplicate / 2nd finding(s) (Single Postings Mode).")

    cp136_c = stats.get('cp136_count', 0)
    cp137_c = stats.get('cp137_count', 0)
    _log(f"Single postings ready for export: {len(records)} (CP-136: {cp136_c}, CP-137: {cp137_c}) - {stats.get('posted_records', 0)} POSTED, {stats.get('not_posted_records', 0)} NOT POSTED")

    # Resolve output destination
    if not output_path or not output_path.strip():
        first_dir = os.path.dirname(file_paths[0])
        first_base = os.path.splitext(os.path.basename(file_paths[0]))[0]
        suffix = "_Single_Postings.xlsx" if deduplicate else "_Extracted.xlsx"
        output_path = os.path.join(first_dir, f"{first_base}{suffix}")
    elif os.path.isdir(output_path.strip()):
        first_base = os.path.splitext(os.path.basename(file_paths[0]))[0]
        suffix = "_Single_Postings.xlsx" if deduplicate else "_Extracted.xlsx"
        output_path = os.path.join(output_path.strip(), f"{first_base}{suffix}")
    elif not output_path.strip().lower().endswith('.xlsx'):
        output_path = output_path.strip() + '.xlsx'
    else:
        output_path = output_path.strip()

    saved_path = export_to_excel(
        records=records,
        output_path=output_path,
        only_present_columns=only_present,
        include_status=include_status,
        include_filename=include_filename or len(file_paths) > 1,
        summary_stats=stats
    )

    _log(f"Successfully saved Excel file: {os.path.basename(saved_path)}")
    return saved_path, stats


# Convenience alias for CP136 & 137
run_cp136_137_check = run_cp136_check
