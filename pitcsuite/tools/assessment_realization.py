"""
Assessment & Realization Extractor Engine & GUI for PITC Suite
Parses WAPDA Assessment & Realization reports for Government Departments (.txt and .pdf),
extracts Sdiv, 14-digit Reference No, Status, Name & Address, Tariff, Department Code,
Arrear Age, OP. Balance, Payments Normal, Adjustor, Write Off, and Net Arrear,
and exports professionally formatted Excel (.xlsx) workbooks matching the official template.
"""

import os
import re
import glob
from pathlib import Path
from datetime import datetime
from typing import List, Dict, Any, Tuple, Optional

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from PySide6.QtCore import Qt, QThread, Signal
from PySide6.QtWidgets import (
    QWidget, QVBoxLayout, QHBoxLayout, QLabel, QPushButton, QLineEdit,
    QListWidget, QGroupBox, QComboBox, QCheckBox, QTextEdit, QProgressBar,
    QFileDialog, QMessageBox, QFrame
)

from pitcsuite.ui_helpers import START_BTN, STOP_BTN, hline, lbl, notify


def parse_amount(val_str: Any) -> Any:
    """
    Parses a monetary or numeric amount string into an int or float.
    Cleans out commas and whitespace.
    If the value has no decimal fraction (e.g. '486145.00'), returns int (486145).
    If it has decimals (e.g. '64.15'), returns float (64.15).
    Returns 0 if blank or unparseable.
    """
    if val_str is None:
        return 0
    if isinstance(val_str, (int, float)):
        return int(val_str) if float(val_str).is_integer() else round(float(val_str), 2)

    s = str(val_str).replace(',', '').strip()
    if not s or s in ('-', '--', 'N/A', 'None'):
        return 0

    # Handle trailing negative sign e.g. '100.00-' or parenthesis '(100.00)'
    is_neg = False
    if s.endswith('-'):
        is_neg = True
        s = s[:-1].strip()
    elif s.startswith('(') and s.endswith(')'):
        is_neg = True
        s = s[1:-1].strip()

    try:
        f = float(s)
        if is_neg:
            f = -f
        return int(f) if f.is_integer() else round(f, 2)
    except ValueError:
        return 0


def parse_assessment_realization_file(filepath: str) -> List[Dict[str, Any]]:
    """
    Parse a single Assessment & Realization report text (.txt) or PDF (.pdf) file.
    Extracts consumer records structured across 3 rows (and optional 4th status row).
    Returns list of dicts with keys:
      - 'Sdiv': 5-digit sub-division string
      - 'Reference No': 14-digit reference number
      - 'Status': string (e.g. 'P-DISC') or None
      - 'Name & Address': combined string from all 3 lines
      - 'Tariff': tariff string (e.g. '66')
      - 'Department Code': integer or string (e.g. 11022)
      - 'Arrear Age': integer arrear age in months
      - 'OP. Balance': opening balance amount
      - 'Payments Normal': payments normal amount
      - 'Adjustor': adjustor amount
      - 'Write Off': write off amount
      - 'Net Arrear': net arrear amount
      - 'Source_File': source file basename
    """
    records = []
    if not os.path.isfile(filepath):
        if os.path.isfile(filepath + ".txt"):
            filepath = filepath + ".txt"
        elif os.path.isfile(filepath + ".pdf"):
            filepath = filepath + ".pdf"
        else:
            return records

    filename = os.path.basename(filepath)
    ext = os.path.splitext(filepath)[1].lower()

    lines = []
    if ext == ".pdf":
        try:
            import pdfplumber
            with pdfplumber.open(filepath) as pdf:
                for page in pdf.pages:
                    txt = page.extract_text(layout=True) or ""
                    lines.extend(txt.splitlines(keepends=True))
        except Exception:
            try:
                import pypdf
                reader = pypdf.PdfReader(filepath)
                for page in reader.pages:
                    txt = page.extract_text() or ""
                    lines.extend(txt.splitlines(keepends=True))
            except Exception:
                lines = []
    else:
        with open(filepath, "r", encoding="utf-8", errors="ignore") as f:
            lines = f.readlines()

    n = len(lines)
    i = 0
    while i < n:
        line = lines[i]

        # Check if line begins with 1-2 digits (batch code / circle)
        m1 = re.match(r'^\s*(\d{1,2})\s+', line)
        if m1 and i + 2 < n:
            l2 = lines[i + 1]
            l3 = lines[i + 2]

            # Line 2 must begin with 5 digits (Subdivision code)
            m2 = re.match(r'^\s*(\d{5})\s+', l2)
            # Line 3 must begin with 7 digits (Account / Consumer number)
            m3 = re.match(r'^\s*(\d{7})\s+', l3)

            if m2 and m3:
                ref1 = m1.group(1).zfill(2)
                ref2 = m2.group(1)
                ref3 = m3.group(1)
                ref_no = f"{ref1}{ref2}{ref3}"
                sdiv = ref2

                # Check for optional 4th line containing Status (e.g. 'P-DISC', 'DISC')
                status = None
                consumed_lines = 3
                if i + 3 < n:
                    l4 = lines[i + 3].strip()
                    if (
                        l4
                        and not l4.startswith('(AGENCY)')
                        and not l4.startswith('(ALL)')
                        and not l4.startswith('---')
                        and not re.match(r'^\d{1,2}\s+', lines[i + 3])
                        and 'REPORT' not in l4
                        and 'DIVISION' not in l4
                        and 'PAGE NO' not in l4
                        and 'REFERENCE NO' not in l4
                    ):
                        # Line 4 contains status e.g. 'P-DISC'
                        status_token = l4.split()[0]
                        # Only take standard status-like tokens (no purely numeric or page artefacts)
                        if any(c.isalpha() for c in status_token):
                            status = status_token
                            consumed_lines = 4

                # Extract Name & Address across the 3 lines (columns 15..45)
                n1 = line[15:45].strip()
                n2 = l2[15:41].strip()
                n3 = l3[15:44].strip()
                name_parts = [p for p in [n1, n2, n3] if p]
                name_addr = " ".join(name_parts)

                # Tariff from Line 1 (columns 45..55)
                tariff = line[45:55].strip()

                # Department Code from Line 2 (columns 41..55)
                dept_raw = l2[41:55].strip().replace(" ", "")
                if dept_raw.isdigit():
                    dept_code = int(dept_raw)
                else:
                    dept_code = dept_raw if dept_raw else None

                # Arrear Age from Line 3 (columns 44..55)
                arr_raw = l3[44:55].strip()
                if arr_raw.isdigit():
                    arr_age = int(arr_raw)
                else:
                    arr_age = arr_raw if arr_raw else 0

                # Numeric tokens from remainder of lines
                tok1 = line[55:].split()
                tok2 = l2[55:].split()
                tok3 = l3[55:].split()

                # OP. Balance (Line 1, column 70..86 or tok1[1])
                op_bal_slice = line[70:86].strip()
                if op_bal_slice:
                    op_balance = parse_amount(op_bal_slice)
                elif len(tok1) >= 2:
                    op_balance = parse_amount(tok1[1])
                else:
                    op_balance = 0

                # Payments Normal (Line 1, column 86..102 or tok1[2])
                pay_norm_slice = line[86:102].strip()
                if pay_norm_slice:
                    payments_normal = parse_amount(pay_norm_slice)
                elif len(tok1) >= 3:
                    payments_normal = parse_amount(tok1[2])
                else:
                    payments_normal = 0

                # Adjustor (Line 2, column 86..102 or tok2[2])
                adj_slice = l2[86:102].strip()
                if adj_slice:
                    adjustor = parse_amount(adj_slice)
                elif len(tok2) >= 3:
                    adjustor = parse_amount(tok2[2])
                else:
                    adjustor = 0

                # Write Off (Line 3, column 86..102 or tok3[2])
                wo_slice = l3[86:102].strip()
                if wo_slice:
                    write_off = parse_amount(wo_slice)
                elif len(tok3) >= 3:
                    write_off = parse_amount(tok3[2])
                else:
                    write_off = 0

                # Net Arrear (Line 3, column 118..134 or tok3[-1])
                net_arr_slice = l3[118:134].strip()
                if net_arr_slice:
                    net_arrear = parse_amount(net_arr_slice)
                elif len(tok3) >= 5:
                    net_arrear = parse_amount(tok3[-1])
                else:
                    net_arrear = 0

                records.append({
                    "Sdiv": sdiv,
                    "Reference No": ref_no,
                    "Status": status,
                    "Name & Address": name_addr,
                    "Tariff": tariff,
                    "Department Code": dept_code,
                    "Arrear Age": arr_age,
                    "OP. Balance": op_balance,
                    "Payments Normal": payments_normal,
                    "Adjustor": adjustor,
                    "Write Off": write_off,
                    "Net Arrear": net_arrear,
                    "Source_File": filename,
                })
                i += consumed_lines
                continue

        i += 1

    return records


def deduplicate_records(records: List[Dict[str, Any]]) -> Tuple[List[Dict[str, Any]], int]:
    """
    Deduplicates records by 'Reference No', keeping the first occurrence.
    Returns (deduped_records, duplicates_removed_count).
    """
    seen = set()
    deduped = []
    dups_count = 0

    for r in records:
        ref = r.get("Reference No")
        if ref and ref in seen:
            dups_count += 1
        else:
            if ref:
                seen.add(ref)
            deduped.append(r)

    return deduped, dups_count


def filter_records(records: List[Dict[str, Any]], status_filter: str = "all") -> List[Dict[str, Any]]:
    """
    Filter records based on status:
      - 'all': return all records
      - 'p_disc': records with Status == 'P-DISC'
      - 'active': records where Status is None or empty
    """
    if status_filter == "p_disc":
        return [r for r in records if r.get("Status") == "P-DISC"]
    elif status_filter == "active":
        return [r for r in records if not r.get("Status")]
    return records


def export_assessment_realization_to_excel(
    records: List[Dict[str, Any]],
    output_path: str,
    include_source_file: bool = False
) -> Dict[str, Any]:
    """
    Exports parsed Assessment & Realization records to an Excel file (.xlsx)
    matching the exact column structure and formats of the reference template:
    Sdiv, Reference No, Status, Name & Address, Tariff, Department Code, Arrear Age,
    OP. Balance, Payments Normal, Adjustor, Write Off, Net Arrear.
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Assessment_Realization"

    # Styling elements
    header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
    header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
    data_font = Font(name="Segoe UI", size=10)
    total_font = Font(name="Segoe UI", size=10, bold=True)
    p_disc_fill = PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid")  # soft red
    total_fill = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")

    thin_border = Border(
        left=Side(style="thin", color="CBD5E1"),
        right=Side(style="thin", color="CBD5E1"),
        top=Side(style="thin", color="CBD5E1"),
        bottom=Side(style="thin", color="CBD5E1")
    )
    total_border = Border(
        left=Side(style="thin", color="CBD5E1"),
        right=Side(style="thin", color="CBD5E1"),
        top=Side(style="thin", color="475569"),
        bottom=Side(style="double", color="475569")
    )

    headers = [
        "Sdiv", "Reference No", "Status", "Name & Address", "Tariff",
        "Department Code", "Arrear Age", "OP. Balance", "Payments Normal",
        "Adjustor", "Write Off", "Net Arrear"
    ]
    if include_source_file:
        headers.append("Source File")

    ws.append(headers)

    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = thin_border
    ws.row_dimensions[1].height = 26

    p_disc_count = 0
    active_count = 0
    tot_op_balance = 0.0
    tot_payments_normal = 0.0
    tot_adjustor = 0.0
    tot_write_off = 0.0
    tot_net_arrear = 0.0

    for r_idx, rec in enumerate(records, start=2):
        # 1. Sdiv (Col A) - text format
        c_sdiv = ws.cell(row=r_idx, column=1, value=str(rec["Sdiv"]))
        c_sdiv.number_format = "@"
        c_sdiv.font = data_font
        c_sdiv.alignment = Alignment(horizontal="center", vertical="center")
        c_sdiv.border = thin_border

        # 2. Reference No (Col B) - text format
        c_ref = ws.cell(row=r_idx, column=2, value=str(rec["Reference No"]))
        c_ref.number_format = "@"
        c_ref.font = data_font
        c_ref.alignment = Alignment(horizontal="center", vertical="center")
        c_ref.border = thin_border

        # 3. Status (Col C)
        status_val = rec.get("Status")
        c_status = ws.cell(row=r_idx, column=3, value=status_val if status_val else None)
        c_status.font = data_font
        c_status.alignment = Alignment(horizontal="center", vertical="center")
        c_status.border = thin_border
        if status_val == "P-DISC":
            c_status.fill = p_disc_fill
            p_disc_count += 1
        else:
            active_count += 1

        # 4. Name & Address (Col D)
        c_name = ws.cell(row=r_idx, column=4, value=rec.get("Name & Address", ""))
        c_name.font = data_font
        c_name.alignment = Alignment(horizontal="left", vertical="center")
        c_name.border = thin_border

        # 5. Tariff (Col E) - text format
        c_tariff = ws.cell(row=r_idx, column=5, value=str(rec.get("Tariff", "")))
        c_tariff.number_format = "@"
        c_tariff.font = data_font
        c_tariff.alignment = Alignment(horizontal="center", vertical="center")
        c_tariff.border = thin_border

        # 6. Department Code (Col F) - integer/code
        dept_val = rec.get("Department Code")
        c_dept = ws.cell(row=r_idx, column=6, value=dept_val)
        c_dept.font = data_font
        c_dept.alignment = Alignment(horizontal="center", vertical="center")
        c_dept.border = thin_border

        # 7. Arrear Age (Col G) - integer
        arr_val = rec.get("Arrear Age")
        c_arr = ws.cell(row=r_idx, column=7, value=arr_val)
        c_arr.font = data_font
        c_arr.alignment = Alignment(horizontal="center", vertical="center")
        c_arr.border = thin_border

        # 8. OP. Balance (Col H) - numeric
        op_val = rec.get("OP. Balance", 0)
        c_op = ws.cell(row=r_idx, column=8, value=op_val)
        c_op.font = data_font
        c_op.alignment = Alignment(horizontal="right", vertical="center")
        c_op.border = thin_border
        if isinstance(op_val, (int, float)):
            tot_op_balance += op_val
            c_op.number_format = "#,##0.00" if isinstance(op_val, float) and not float(op_val).is_integer() else "#,##0"

        # 9. Payments Normal (Col I) - numeric
        pn_val = rec.get("Payments Normal", 0)
        c_pn = ws.cell(row=r_idx, column=9, value=pn_val)
        c_pn.font = data_font
        c_pn.alignment = Alignment(horizontal="right", vertical="center")
        c_pn.border = thin_border
        if isinstance(pn_val, (int, float)):
            tot_payments_normal += pn_val
            c_pn.number_format = "#,##0.00" if isinstance(pn_val, float) and not float(pn_val).is_integer() else "#,##0"

        # 10. Adjustor (Col J) - numeric
        adj_val = rec.get("Adjustor", 0)
        c_adj = ws.cell(row=r_idx, column=10, value=adj_val)
        c_adj.font = data_font
        c_adj.alignment = Alignment(horizontal="right", vertical="center")
        c_adj.border = thin_border
        if isinstance(adj_val, (int, float)):
            tot_adjustor += adj_val
            c_adj.number_format = "#,##0.00" if isinstance(adj_val, float) and not float(adj_val).is_integer() else "#,##0"

        # 11. Write Off (Col K) - numeric
        wo_val = rec.get("Write Off", 0)
        c_wo = ws.cell(row=r_idx, column=11, value=wo_val)
        c_wo.font = data_font
        c_wo.alignment = Alignment(horizontal="right", vertical="center")
        c_wo.border = thin_border
        if isinstance(wo_val, (int, float)):
            tot_write_off += wo_val
            c_wo.number_format = "#,##0.00" if isinstance(wo_val, float) and not float(wo_val).is_integer() else "#,##0"

        # 12. Net Arrear (Col L) - numeric
        na_val = rec.get("Net Arrear", 0)
        c_na = ws.cell(row=r_idx, column=12, value=na_val)
        c_na.font = data_font
        c_na.alignment = Alignment(horizontal="right", vertical="center")
        c_na.border = thin_border
        if isinstance(na_val, (int, float)):
            tot_net_arrear += na_val
            c_na.number_format = "#,##0.00" if isinstance(na_val, float) and not float(na_val).is_integer() else "#,##0"

        # 13. Source File (Col M, optional)
        if include_source_file:
            c_src = ws.cell(row=r_idx, column=13, value=rec.get("Source_File", ""))
            c_src.font = data_font
            c_src.alignment = Alignment(horizontal="left", vertical="center")
            c_src.border = thin_border

        ws.row_dimensions[r_idx].height = 20

    # Summary / Total Row
    tot_row = len(records) + 2
    ws.row_dimensions[tot_row].height = 24

    for c in range(1, len(headers) + 1):
        cell = ws.cell(row=tot_row, column=c)
        cell.fill = total_fill
        cell.border = total_border
        cell.font = total_font

    ws.cell(row=tot_row, column=1, value="TOTAL").alignment = Alignment(horizontal="center", vertical="center")

    # Sum formulas
    last_data_row = tot_row - 1
    if records:
        ws.cell(row=tot_row, column=8, value=f"=SUM(H2:H{last_data_row})").number_format = "#,##0.00" if isinstance(tot_op_balance, float) and not float(tot_op_balance).is_integer() else "#,##0"
        ws.cell(row=tot_row, column=8).alignment = Alignment(horizontal="right", vertical="center")

        ws.cell(row=tot_row, column=9, value=f"=SUM(I2:I{last_data_row})").number_format = "#,##0.00" if isinstance(tot_payments_normal, float) and not float(tot_payments_normal).is_integer() else "#,##0"
        ws.cell(row=tot_row, column=9).alignment = Alignment(horizontal="right", vertical="center")

        ws.cell(row=tot_row, column=10, value=f"=SUM(J2:J{last_data_row})").number_format = "#,##0.00" if isinstance(tot_adjustor, float) and not float(tot_adjustor).is_integer() else "#,##0"
        ws.cell(row=tot_row, column=10).alignment = Alignment(horizontal="right", vertical="center")

        ws.cell(row=tot_row, column=11, value=f"=SUM(K2:K{last_data_row})").number_format = "#,##0.00" if isinstance(tot_write_off, float) and not float(tot_write_off).is_integer() else "#,##0"
        ws.cell(row=tot_row, column=11).alignment = Alignment(horizontal="right", vertical="center")

        ws.cell(row=tot_row, column=12, value=f"=SUM(L2:L{last_data_row})").number_format = "#,##0.00" if isinstance(tot_net_arrear, float) and not float(tot_net_arrear).is_integer() else "#,##0"
        ws.cell(row=tot_row, column=12).alignment = Alignment(horizontal="right", vertical="center")

    # Column Widths auto-fit with comfortable minimums
    min_widths = {
        1: 10,  # Sdiv
        2: 18,  # Reference No
        3: 12,  # Status
        4: 45,  # Name & Address
        5: 10,  # Tariff
        6: 18,  # Department Code
        7: 12,  # Arrear Age
        8: 15,  # OP. Balance
        9: 17,  # Payments Normal
        10: 13, # Adjustor
        11: 13, # Write Off
        12: 15, # Net Arrear
        13: 25, # Source File
    }

    for col_idx in range(1, len(headers) + 1):
        col_letter = get_column_letter(col_idx)
        max_len = len(str(headers[col_idx - 1]))
        for r in range(2, min(tot_row, 100)):
            val = ws.cell(row=r, column=col_idx).value
            if val is not None:
                max_len = max(max_len, len(str(val)))
        target_width = max(max_len + 3, min_widths.get(col_idx, 12))
        ws.column_dimensions[col_letter].width = target_width

    # Freeze header row
    ws.freeze_panes = "A2"

    os.makedirs(os.path.dirname(os.path.abspath(output_path)), exist_ok=True)
    wb.save(output_path)

    return {
        "total": len(records),
        "p_disc": p_disc_count,
        "active": active_count,
        "total_op_balance": tot_op_balance,
        "total_payments_normal": tot_payments_normal,
        "total_adjustor": tot_adjustor,
        "total_write_off": tot_write_off,
        "total_net_arrear": tot_net_arrear,
    }


class AssessmentRealizationWorker(QThread):
    """Worker thread for parsing Assessment & Realization text file asynchronously."""
    log = Signal(str)
    progress = Signal(int, int, str)
    done = Signal(str, dict, list)
    failed = Signal(str)

    def __init__(self, file_path: str, output_path: str):
        super().__init__()
        self.file_path = file_path
        self.output_path = output_path

    def run(self):
        try:
            fname = os.path.basename(self.file_path)
            self.log.emit(f"🚀 Starting extraction from: {fname}")
            self.progress.emit(15, 100, f"Reading {fname}...")

            records = parse_assessment_realization_file(self.file_path)
            self.log.emit(f"   ↳ Extracted {len(records)} consumer record(s) from {fname}")

            if not records:
                raise ValueError("No consumer records found in the specified file.")

            if self.isInterruptionRequested():
                self.log.emit("⚠ Process cancelled by user.")
                return

            self.progress.emit(65, 100, "Generating Excel report...")
            self.log.emit(f"💾 Exporting to Excel: {self.output_path}...")
            stats = export_assessment_realization_to_excel(records, self.output_path)

            self.log.emit("✔ Export completed successfully!")
            self.log.emit(f"   • Total records in Excel: {stats['total']}")
            self.log.emit(f"   • Active / Connected: {stats['active']}")
            self.log.emit(f"   • Disconnected (P-DISC): {stats['p_disc']}")
            self.log.emit(f"   • Total OP. Balance: Rs. {stats['total_op_balance']:,.2f}")
            self.log.emit(f"   • Total Net Arrear: Rs. {stats['total_net_arrear']:,.2f}")
            self.log.emit(f"   • Output file: {self.output_path}")

            self.done.emit(self.output_path, stats, records)
        except Exception as exc:
            self.failed.emit(str(exc))


class AssessmentRealizationPanel(QWidget):
    """UI Panel for Assessment & Realization Extractor in PITC Suite."""
    def __init__(self):
        super().__init__()
        self.worker = None
        self.last_output = ""

        layout = QVBoxLayout(self)
        layout.setSpacing(10)

        # Header Title and Description
        layout.addWidget(lbl("Govt Assessment Realization", bold=True, color="#7eb8f7"))
        layout.addWidget(lbl(
            "Import Assessment & Realization text report (.txt) for Government Departments and export structured Excel sheet.",
            color="#556"
        ))
        layout.addWidget(hline())

        # Inputs Group
        group = QGroupBox("Inputs")
        form = QVBoxLayout(group)

        # Source File row
        r1 = QHBoxLayout()
        r1.addWidget(QLabel("Source Report (.txt):"))
        self.file_ed = QLineEdit()
        self.file_ed.setReadOnly(True)
        self.file_ed.setPlaceholderText("Select Assessment & Realization text file...")
        r1.addWidget(self.file_ed, 1)

        self.add_file_btn = QPushButton("📄 Add File")
        self.add_file_btn.clicked.connect(self._add_file)
        r1.addWidget(self.add_file_btn)
        form.addLayout(r1)

        # Target Excel File row
        r2 = QHBoxLayout()
        r2.addWidget(QLabel("Target Excel File:"))
        self.excel_ed = QLineEdit()
        self.excel_ed.setPlaceholderText("Excel output path (auto-suggested on file selection)...")
        r2.addWidget(self.excel_ed, 1)

        self.browse_excel_btn = QPushButton("📁 Browse...")
        self.browse_excel_btn.clicked.connect(self._browse_excel)
        r2.addWidget(self.browse_excel_btn)
        form.addLayout(r2)

        layout.addWidget(group)

        # Actions Row
        actions = QHBoxLayout()
        self.start_btn = QPushButton("▶  START EXTRACTION")
        self.start_btn.setStyleSheet(START_BTN)
        self.start_btn.clicked.connect(self._start)
        actions.addWidget(self.start_btn)

        self.stop_btn = QPushButton("■  STOP")
        self.stop_btn.setStyleSheet(STOP_BTN)
        self.stop_btn.setEnabled(False)
        self.stop_btn.clicked.connect(self._stop)
        actions.addWidget(self.stop_btn)

        self.open_excel_btn = QPushButton("📊 Open Excel")
        self.open_excel_btn.setEnabled(False)
        self.open_excel_btn.clicked.connect(self._open_excel)
        actions.addWidget(self.open_excel_btn)

        self.open_folder_btn = QPushButton("📂 Open Folder")
        self.open_folder_btn.setEnabled(False)
        self.open_folder_btn.clicked.connect(self._open_folder)
        actions.addWidget(self.open_folder_btn)
        actions.addStretch()
        layout.addLayout(actions)

        # Status & Progress
        self.status = QLabel("Status: Ready. Please select an Assessment Realization text file.")
        self.status.setStyleSheet("color:#4a90d9;")
        layout.addWidget(self.status)

        self.progress_bar = QProgressBar()
        self.progress_bar.setRange(0, 100)
        self.progress_bar.setValue(0)
        layout.addWidget(self.progress_bar)

        # Log Box
        self.log_box = QTextEdit()
        self.log_box.setReadOnly(True)
        layout.addWidget(self.log_box)

    def _add_file(self):
        path, _ = QFileDialog.getOpenFileName(
            self, "Select Assessment & Realization Report", "",
            "Text Files (*.txt);;All Files (*.*)"
        )
        if path:
            self.file_ed.setText(path)
            folder = os.path.dirname(path)
            base = os.path.splitext(os.path.basename(path))[0]
            date_str = datetime.now().strftime("%d%m%Y")
            self.excel_ed.setText(os.path.join(folder, f"{base}_Assessment_Realization_{date_str}.xlsx"))

    def _browse_excel(self):
        path, _ = QFileDialog.getSaveFileName(
            self, "Save Assessment Realization Excel File", "", "Excel Workbook (*.xlsx);;All Files (*.*)"
        )
        if path:
            if not path.lower().endswith(".xlsx"):
                path += ".xlsx"
            self.excel_ed.setText(path)

    def _start(self):
        src_path = self.file_ed.text().strip()
        if not src_path or not os.path.exists(src_path):
            QMessageBox.warning(self, "Missing File", "Please select a valid Assessment & Realization text file.")
            return

        out_path = self.excel_ed.text().strip()
        if not out_path:
            folder = os.path.dirname(src_path)
            base = os.path.splitext(os.path.basename(src_path))[0]
            date_str = datetime.now().strftime("%d%m%Y")
            out_path = os.path.join(folder, f"{base}_Assessment_Realization_{date_str}.xlsx")
            self.excel_ed.setText(out_path)

        if not out_path.lower().endswith(".xlsx"):
            out_path += ".xlsx"

        self.last_output = out_path
        self.log_box.clear()
        self.progress_bar.setValue(0)
        self.status.setText("Status: Extracting data from report...")
        self.start_btn.setEnabled(False)
        self.stop_btn.setEnabled(True)
        self.open_excel_btn.setEnabled(False)
        self.open_folder_btn.setEnabled(False)

        self.worker = AssessmentRealizationWorker(file_path=src_path, output_path=out_path)
        self.worker.log.connect(self.log_box.append)
        self.worker.progress.connect(self._progress)
        self.worker.done.connect(self._done)
        self.worker.failed.connect(self._failed)
        self.worker.finished.connect(lambda: (
            self.start_btn.setEnabled(True),
            self.stop_btn.setEnabled(False)
        ))
        self.worker.start()

    def _stop(self):
        if self.worker and self.worker.isRunning():
            self.worker.requestInterruption()
            self.status.setText("Status: Stopping extraction...")

    def _progress(self, current: int, total: int, msg: str):
        pct = int((current / max(1, total)) * 100)
        self.progress_bar.setValue(pct)
        self.status.setText(f"Status: {msg}")

    def _done(self, out_file: str, stats: dict, records: list):
        self.progress_bar.setValue(100)
        self.status.setText("Status: Extraction completed successfully!")
        self.open_excel_btn.setEnabled(True)
        self.open_folder_btn.setEnabled(True)

        notify(
            self,
            "Extraction Complete",
            f"Successfully extracted {stats.get('total', 0)} record(s) to Excel:\n{out_file}",
            file_path=out_file
        )

    def _failed(self, err_msg: str):
        self.progress_bar.setValue(0)
        self.status.setText("Status: Extraction failed.")
        notify(self, "Error", f"Assessment Realization extraction failed:\n{err_msg}", critical=True)

    def _open_excel(self):
        if self.last_output and os.path.exists(self.last_output):
            os.startfile(self.last_output)

    def _open_folder(self):
        if self.last_output:
            folder = os.path.dirname(os.path.abspath(self.last_output))
            if os.path.exists(folder):
                os.startfile(folder)

