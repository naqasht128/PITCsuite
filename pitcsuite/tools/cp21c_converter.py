import os
import re
import fitz
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

COL_HEADERS = [
    'Sr No.',
    'TARIFF CODE',
    'PREV. METER STATUS',
    'SANC. LOAD',
    'REFERENCE NUMBER',
    'M.F',
    'METER SEQ',
    'PREVIOUS KWH METER READING',
    'RANGE',
    'PRESENT KWH METER READING',
    'UNITS TO BE BILLED',
    'MDI Reading',
    'METER NO',
    'STATUS CODE',
    'CORRS. MONTH UNITS',
    'FEEDER CD',
    'SUB-BATCH',
    'METER READER NAME'
]

def assign_col(w):
    """Maps a word span to its appropriate table column based on x-coordinate."""
    x0, y0, x1, y1, text = w[0], w[1], w[2], w[3], w[4].strip()
    xc = (x0 + x1) / 2.0
    
    # Special exact match for 14-digit reference number
    if len(text) == 14 and text.isdigit():
        return 'REFERENCE NUMBER'
    
    if xc < 38:
        return 'Sr No.'
    elif 38 <= xc < 58:
        return 'TARIFF CODE'
    elif 58 <= xc < 95:
        return 'PREV. METER STATUS'
    elif 95 <= xc < 138:
        return 'SANC. LOAD'
    elif 138 <= xc < 205:
        return 'REFERENCE NUMBER'
    elif 205 <= xc < 223:
        return 'M.F'
    elif 223 <= xc < 235:
        return 'METER SEQ'
    elif 235 <= xc < 285:
        return 'PREVIOUS KWH METER READING'
    elif 285 <= xc < 302:
        return 'RANGE'
    elif 302 <= xc < 345:
        return 'PRESENT KWH METER READING'
    elif 345 <= xc < 375:
        return 'UNITS TO BE BILLED'
    elif 375 <= xc < 420:
        return 'MDI Reading'
    elif 420 <= xc < 475:
        return 'METER NO'
    elif 475 <= xc < 498:
        return 'STATUS CODE'
    elif 498 <= xc < 545:
        return 'CORRS. MONTH UNITS'
    elif xc >= 545:
        return 'FEEDER CD'
    return 'UNKNOWN'

def convert_cp21c_pdf(pdf_path, excel_path, progress_callback=None, log_callback=None):
    """
    Extracts CP-21C electricity meter reading data from PDF and exports to a single-sheet Excel file.
    """
    def log(msg, level='INFO'):
        if log_callback:
            log_callback(msg)
        else:
            print(f"[{level}] {msg}")

    if not os.path.exists(pdf_path):
        raise FileNotFoundError(f"PDF file not found: {pdf_path}")

    log(f"Opening PDF: {os.path.basename(pdf_path)}", "INFO")
    doc = fitz.open(pdf_path)
    total_pages = len(doc)
    log(f"Total pages to process: {total_pages}", "INFO")

    all_records = []
    current_sub_batch = "01"
    current_meter_reader = ""
    sub_batch_readers = {}
    expected_batch_units = None
    expected_batch_corrs = None

    for pno in range(total_pages):
        page = doc[pno]
        page_num = pno + 1
        
        if progress_callback:
            progress_callback(page_num, total_pages, f"Reading Page {page_num} of {total_pages}...")

        words = page.get_text('words')
        
        sb_words = [w for w in words if 115 <= w[1] <= 140]
        sb_words.sort(key=lambda w: w[0])
        sb_text = ' '.join([w[4] for w in sb_words])
        
        if 'sub-batch:' in sb_text.lower():
            m_sb = re.search(r'sub-batch:\s*(\d+)', sb_text, re.IGNORECASE)
            if m_sb:
                current_sub_batch = f"{int(m_sb.group(1)):02d}"
            m_reader = re.search(r'meter\s*reader\s*name:\s*(.*?)(?:\s+\d+)?$', sb_text, re.IGNORECASE)
            if m_reader:
                reader_raw = m_reader.group(1).strip()
                reader_clean = re.sub(r'\s+\d+$', '', reader_raw).strip()
                current_meter_reader = reader_clean
            else:
                current_meter_reader = ""
            sub_batch_readers[current_sub_batch] = current_meter_reader
            log(f"Page {page_num}: Sub-Batch {current_sub_batch} (Reader: {current_meter_reader or 'N/A'})")

        text_full = page.get_text()
        if 'batch units total' in text_full.lower():
            m_bu = re.search(r'batch\s*units\s*total\s*(\d+)', text_full, re.IGNORECASE)
            if m_bu:
                expected_batch_units = int(m_bu.group(1))
            m_cu = re.search(r'corrs\.\s*month\s*batch\s*units\s*total\s*(\d+)', text_full, re.IGNORECASE)
            if m_cu:
                expected_batch_corrs = int(m_cu.group(1))

        clean_words = []
        for w in words:
            txt = w[4].strip()
            y0 = w[1]
            if y0 < 112 or y0 > 745:
                continue
            if any(k in txt for k in ['Signatures', 'Processed', 'SUPERVISOR', 'SECTION', 'S.D.O', '=====']):
                continue
            clean_words.append(w)

        clean_words.sort(key=lambda w: (w[1], w[0]))
        lines = []
        curr_line = []
        curr_y = None
        for w in clean_words:
            if curr_y is None:
                curr_y = w[1]
                curr_line.append(w)
            elif abs(w[1] - curr_y) <= 4:
                curr_line.append(w)
            else:
                lines.append(curr_line)
                curr_line = [w]
                curr_y = w[1]
        if curr_line:
            lines.append(curr_line)

        page_rows = 0
        for line in lines:
            line_txt = ' '.join([w[4] for w in line])
            words_in_line = [w[4].lower().rstrip(':') for w in line]
            
            if any(k in words_in_line for k in ['sub-batch', 'reader', 'sub', 'batch', 'total', 'corrs.', 'month', 'advance', 'signatures', 'report']):
                continue

            row_dict = {}
            for w in line:
                col = assign_col(w)
                val = w[4].strip()
                if col in row_dict:
                    row_dict[col] += ' ' + val
                else:
                    row_dict[col] = val

            row_dict['page'] = page_num
            row_dict['sub_batch'] = current_sub_batch
            row_dict['meter_reader'] = current_meter_reader
            row_dict['raw_line'] = line_txt
            all_records.append(row_dict)
            page_rows += 1

    log(f"Extracted {len(all_records)} total meter reading records across {total_pages} pages.")
    log(f"Creating Excel workbook...")
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Meter Reading List"
    ws.views.sheetView[0].showGridLines = True

    header_font = Font(name='Calibri', size=11, bold=True, color='000000')
    header_fill = PatternFill(start_color='D9E1F2', end_color='D9E1F2', fill_type='solid')
    header_align = Alignment(horizontal='center', vertical='center', wrap_text=True)
    
    thin_border = Border(
        left=Side(style='thin', color='D3D3D3'),
        right=Side(style='thin', color='D3D3D3'),
        top=Side(style='thin', color='D3D3D3'),
        bottom=Side(style='thin', color='D3D3D3')
    )

    for col_idx, h_name in enumerate(COL_HEADERS, start=1):
        cell = ws.cell(row=1, column=col_idx, value=h_name)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_align
        cell.border = thin_border
    ws.row_dimensions[1].height = 28

    current_parent = None
    total_units_billed = 0
    total_corrs_units = 0

    for row_idx, r in enumerate(all_records, start=2):
        is_seq2 = (r.get('METER SEQ') == '2')
        if not is_seq2:
            current_parent = r

        sr_val = r.get('Sr No.') or (current_parent.get('Sr No.') if is_seq2 and current_parent else None)
        tariff_val = r.get('TARIFF CODE') or (current_parent.get('TARIFF CODE') if is_seq2 and current_parent else '')
        prev_status_val = r.get('PREV. METER STATUS', '')
        sanc_val = r.get('SANC. LOAD') or (current_parent.get('SANC. LOAD') if is_seq2 and current_parent else None)
        ref_val = r.get('REFERENCE NUMBER') or (current_parent.get('REFERENCE NUMBER') if is_seq2 and current_parent else '')
        mf_val = r.get('M.F', '')
        seq_val = r.get('METER SEQ', '1')
        prev_rd_val = r.get('PREVIOUS KWH METER READING')
        range_val = r.get('RANGE')
        pres_rd_val = r.get('PRESENT KWH METER READING')
        units_val = r.get('UNITS TO BE BILLED')
        mdi_val = r.get('MDI Reading', '')
        meter_val = r.get('METER NO') or (current_parent.get('METER NO') if is_seq2 and current_parent else '')
        status_val = r.get('STATUS CODE', '')
        corrs_val = r.get('CORRS. MONTH UNITS')
        feeder_val = r.get('FEEDER CD') or (current_parent.get('FEEDER CD') if is_seq2 and current_parent else '')
        sb_val = r.get('sub_batch', '01')
        reader_val = r.get('meter_reader', '')

        c1 = ws.cell(row=row_idx, column=1)
        if sr_val is not None and str(sr_val).strip().isdigit():
            c1.value = int(sr_val)
            c1.number_format = '0'
            c1.alignment = Alignment(horizontal='center')
        c1.border = thin_border

        c2 = ws.cell(row=row_idx, column=2)
        if tariff_val:
            c2.value = str(tariff_val).strip()
            c2.number_format = '@'
            c2.alignment = Alignment(horizontal='center')
        c2.border = thin_border

        c3 = ws.cell(row=row_idx, column=3)
        if prev_status_val:
            c3.value = str(prev_status_val).strip()
            c3.number_format = '@'
            c3.alignment = Alignment(horizontal='center')
        c3.border = thin_border

        c4 = ws.cell(row=row_idx, column=4)
        if sanc_val is not None and str(sanc_val).strip():
            try:
                c4.value = float(sanc_val)
                c4.number_format = '0.00'
                c4.alignment = Alignment(horizontal='right')
            except ValueError:
                c4.value = str(sanc_val).strip()
        c4.border = thin_border

        c5 = ws.cell(row=row_idx, column=5)
        if ref_val:
            c5.value = str(ref_val).strip()
            c5.number_format = '@'
            c5.alignment = Alignment(horizontal='center')
        c5.border = thin_border

        c6 = ws.cell(row=row_idx, column=6)
        if mf_val is not None and str(mf_val).strip():
            try:
                c6.value = float(mf_val)
                c6.number_format = '0.00'
            except ValueError:
                c6.value = str(mf_val).strip()
            c6.alignment = Alignment(horizontal='right')
        c6.border = thin_border

        c7 = ws.cell(row=row_idx, column=7)
        if seq_val is not None and str(seq_val).strip().isdigit():
            c7.value = int(seq_val)
            c7.number_format = '0'
            c7.alignment = Alignment(horizontal='center')
        c7.border = thin_border

        c8 = ws.cell(row=row_idx, column=8)
        if prev_rd_val is not None and str(prev_rd_val).strip():
            try:
                c8.value = int(round(float(prev_rd_val)))
                c8.number_format = '0'
                c8.alignment = Alignment(horizontal='right')
            except ValueError:
                pass
        c8.border = thin_border

        c9 = ws.cell(row=row_idx, column=9)
        if range_val is not None and str(range_val).strip().isdigit():
            c9.value = int(range_val)
            c9.number_format = '0'
            c9.alignment = Alignment(horizontal='center')
        c9.border = thin_border

        c10 = ws.cell(row=row_idx, column=10)
        if pres_rd_val is not None and str(pres_rd_val).strip():
            try:
                c10.value = int(round(float(pres_rd_val)))
                c10.number_format = '0'
                c10.alignment = Alignment(horizontal='right')
            except ValueError:
                pass
        c10.border = thin_border

        c11 = ws.cell(row=row_idx, column=11)
        if units_val is not None and str(units_val).strip():
            try:
                u_int = int(round(float(units_val)))
                c11.value = u_int
                c11.number_format = '0'
                c11.alignment = Alignment(horizontal='right')
                total_units_billed += u_int
            except ValueError:
                pass
        c11.border = thin_border

        c12 = ws.cell(row=row_idx, column=12)
        if mdi_val is not None and str(mdi_val).strip():
            try:
                c12.value = float(mdi_val)
                c12.number_format = '0.00'
                c12.alignment = Alignment(horizontal='right')
            except ValueError:
                c12.value = str(mdi_val).strip()
        c12.border = thin_border

        c13 = ws.cell(row=row_idx, column=13)
        if meter_val:
            c13.value = str(meter_val).strip()
            c13.number_format = '@'
            c13.alignment = Alignment(horizontal='center')
        c13.border = thin_border

        c14 = ws.cell(row=row_idx, column=14)
        if status_val is not None and str(status_val).strip():
            st_clean = str(status_val).strip()
            if st_clean.isdigit():
                c14.value = int(st_clean)
                c14.number_format = '0'
            else:
                c14.value = st_clean
                c14.number_format = '@'
            c14.alignment = Alignment(horizontal='center')
        c14.border = thin_border

        c15 = ws.cell(row=row_idx, column=15)
        if corrs_val is not None and str(corrs_val).strip():
            try:
                c_int = int(round(float(corrs_val)))
                c15.value = c_int
                c15.number_format = '0'
                c15.alignment = Alignment(horizontal='right')
                total_corrs_units += c_int
            except ValueError:
                pass
        c15.border = thin_border

        c16 = ws.cell(row=row_idx, column=16)
        if feeder_val is not None and str(feeder_val).strip():
            fd_clean = str(feeder_val).strip()
            if fd_clean.isdigit():
                c16.value = int(fd_clean)
                c16.number_format = '0'
            else:
                c16.value = fd_clean
                c16.number_format = '@'
            c16.alignment = Alignment(horizontal='center')
        c16.border = thin_border

        c17 = ws.cell(row=row_idx, column=17)
        c17.value = sb_val
        c17.number_format = '@'
        c17.alignment = Alignment(horizontal='center')
        c17.border = thin_border

        c18 = ws.cell(row=row_idx, column=18)
        c18.value = reader_val
        c18.number_format = '@'
        c18.alignment = Alignment(horizontal='left')
        c18.border = thin_border

    col_widths = {
        'A': 10, 'B': 13, 'C': 18, 'D': 13, 'E': 20,
        'F': 10, 'G': 12, 'H': 18, 'I': 10, 'J': 18,
        'K': 16, 'L': 14, 'M': 16, 'N': 14, 'O': 16,
        'P': 13, 'Q': 13, 'R': 22
    }
    for col_letter, width in col_widths.items():
        ws.column_dimensions[col_letter].width = width

    os.makedirs(os.path.dirname(os.path.abspath(excel_path)), exist_ok=True)
    wb.save(excel_path)
    log(f"Excel workbook saved successfully to: {excel_path}")
    
    log(f"--- Summary & Reconciliation ---")
    log(f"Total Rows Written: {len(all_records)}")
    log(f"Total Units Billed: {total_units_billed:,}")
    log(f"Total Corrs. Month Units: {total_corrs_units:,}")
    if expected_batch_units is not None:
        diff_units = total_units_billed - expected_batch_units
        log(f"Expected Batch Units Total: {expected_batch_units:,} (Difference: {diff_units})")
    if expected_batch_corrs is not None:
        diff_corrs = total_corrs_units - expected_batch_corrs
        log(f"Expected Corrs. Units Total: {expected_batch_corrs:,} (Difference: {diff_corrs})")

    return {
        'total_rows': len(all_records),
        'total_units': total_units_billed,
        'total_corrs': total_corrs_units,
        'expected_units': expected_batch_units,
        'expected_corrs': expected_batch_corrs,
        'sub_batches': sub_batch_readers
    }
